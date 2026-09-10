# -*- coding: utf-8 -*-
"""FinanceSystem Windows desktop UI.

The GUI deliberately owns no accounting state transitions.  It validates user
input, renders read models, and delegates every voucher/data-issue/period write
to the service layer.  Keeping that boundary explicit makes the same UI usable
from a source checkout and from the frozen Windows build.
"""
from __future__ import annotations

import csv
import datetime as dt
import inspect
import json
import os
from pathlib import Path
import re
import sys
import threading
import tkinter as tk
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from tkinter import filedialog, messagebox, simpledialog, ttk
import uuid

import db
from modules import arap, expense as expense_mod, feishu, invoice_gui, ledger, reports
from gui.ui_theme import (
    C_BG, C_CARD, C_CARD_HOVER, C_BORDER, C_NAV, C_NAV_ACTIVE,
    C_NAV_HOVER, C_NAV_TEXT, C_NAV_TEXT_ACTIVE,
    C_TEXT, C_TEXT_MUTED, C_TEXT_DIM,
    C_ACCENT, C_ACCENT_HOVER, C_ACCENT_SOFT,
    C_ROW_ALT,
    C_OK, C_OK_SOFT, C_WARN, C_WARN_SOFT, C_DANGER, C_DANGER_SOFT,
    C_INPUT_BG, C_INPUT_FG, C_INPUT_BORDER,
    FONT, FONT_SMALL, FONT_B, FONT_TITLE, FONT_H2, FONT_H3, FONT_NUM, FONT_MONO,
    configure_dark_style,
    CardFrame,
    ModernButton,
    NavButton,
    StatCard,
    C_PANEL,
)

try:
    import app_paths
except Exception:
    app_paths = None
try:
    import security
except Exception:
    security = None

ROOT = Path(getattr(app_paths, "INSTALL_ROOT", Path(__file__).resolve().parent.parent))


def _enable_dpi_awareness() -> None:
    """Enable per-monitor DPI before Tk creates a native window."""
    if sys.platform != "win32":
        return
    try:
        import ctypes

        ctypes.windll.shcore.SetProcessDpiAwareness(2)
    except Exception:
        try:
            ctypes.windll.user32.SetProcessDPIAware()
        except Exception:
            pass



STATUS_LABELS = {
    "draft": "草稿",
    "submitted": "待审核",
    "pending_review": "待审核",
    "pending_approval": "待审批",
    "audited": "已审核",
    "approved": "已批准",
    "posted": "已记账",
    "reversed": "已冲销",
    "rejected": "已驳回",
    "accounted": "已记账",
    "paid": "已付款",
    "confirmed": "已确认",
    "pending_confirmation": "待确认",
}

VOUCHER_STATUS_CODES = {
    "草稿": "draft",
    "待审核": "pending_review",
    "已审核": "audited",
    "已记账": "posted",
    "已冲销": "reversed",
}


def _today() -> str:
    return dt.date.today().isoformat()


def _latest_posted_period() -> str | None:
    """最近一个有「已记账」凭证的会计期间；无则返回 None。"""
    try:
        conn = db.get_conn()
        try:
            row = conn.execute(
                "SELECT period FROM voucher WHERE status='posted' AND period IS NOT NULL "
                "ORDER BY period DESC LIMIT 1"
            ).fetchone()
        finally:
            conn.close()
        return str(row[0]) if row and row[0] else None
    except Exception:
        return None


_ALL_ACCOUNTS = "全部科目"


def _current_period() -> str:
    return dt.date.today().strftime("%Y-%m")


def _validate_date(value: str) -> str:
    try:
        return dt.date.fromisoformat(value.strip()).isoformat()
    except (TypeError, ValueError) as exc:
        raise ValueError("日期必须是有效的 YYYY-MM-DD") from exc


def _validate_period(value: str) -> str:
    value = value.strip()
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", value):
        raise ValueError("期间必须是 YYYY-MM，例如 2026-07")
    return value


def _parse_cents(value: object, *, positive: bool = False) -> int:
    text = str(value or "").replace(",", "").strip()
    try:
        cents = int((Decimal(text or "0") * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    except (InvalidOperation, ValueError) as exc:
        raise ValueError("金额格式无效") from exc
    if positive and cents <= 0:
        raise ValueError("金额必须大于 0")
    if cents < 0:
        raise ValueError("金额不能为负数")
    return cents


def _money(cents: object) -> str:
    try:
        value = int(cents or 0)
    except (TypeError, ValueError):
        value = 0
    return f"{value / 100:,.2f}"


def _row_cents(row: dict, cents_key: str, legacy_key: str) -> int:
    if row.get(cents_key) is not None:
        try:
            return int(row[cents_key])
        except (TypeError, ValueError):
            pass
    try:
        return int((Decimal(str(row.get(legacy_key) or 0)) * 100).quantize(Decimal("1")))
    except (InvalidOperation, ValueError):
        return 0


def _status(value: object) -> str:
    text = str(value or "")
    return STATUS_LABELS.get(text, text)


def _result(value) -> tuple[bool, str, object]:
    """Normalize tuples and IntegrationResult objects for UI messages."""
    if hasattr(value, "ok"):
        return bool(value.ok), str(getattr(value, "message", "")), getattr(value, "data", None)
    if isinstance(value, tuple):
        items = list(value) + [None, None, None]
        return bool(items[0]), str(items[1] or ""), items[2]
    if isinstance(value, bool):
        return value, "操作成功" if value else "操作失败", None
    return True, str(value or "操作成功"), value


def _call_supported(function, **kwargs):
    """Call a service while tolerating a rolling upgrade's optional kwargs."""
    signature = inspect.signature(function)
    if any(p.kind == inspect.Parameter.VAR_KEYWORD for p in signature.parameters.values()):
        return function(**kwargs)
    selected = {key: value for key, value in kwargs.items() if key in signature.parameters}
    return function(**selected)


def _table_exists(conn, table: str) -> bool:
    return bool(
        conn.execute(
            "SELECT 1 FROM sqlite_master WHERE type='table' AND name=?", (table,)
        ).fetchone()
    )


def _table_columns(conn, table: str) -> set[str]:
    if not _table_exists(conn, table):
        return set()
    return {row[1] for row in conn.execute('PRAGMA table_info("%s")' % table)}


class _ZebraTree(ttk.Treeview):
    """Treeview with alternating row backgrounds, keeping the selection visible.

    ttk paints a tag's ``background`` *over* the style's ``selected`` state, so a
    naive zebra tag would make the selected row look unselected.  We therefore
    drive a dedicated ``sel`` tag on every selection change.
    """

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.tag_configure("even", background=C_CARD)
        self.tag_configure("odd", background=C_ROW_ALT)
        self.tag_configure("sel", background=C_ACCENT_SOFT, foreground=C_ACCENT)
        self.bind("<<TreeviewSelect>>", self._refresh_selection, add="+")

    def insert(self, parent, index, *args, **kwargs):
        parity = "odd" if len(self.get_children(parent)) % 2 == 1 else "even"
        tags = [parity, *list(kwargs.get("tags") or ())]
        kwargs["tags"] = tags
        return super().insert(parent, index, *args, **kwargs)

    def _refresh_selection(self, _event=None) -> None:
        selected = set(self.selection())
        for iid in self.get_children(""):
            tags = [t for t in self.item(iid, "tags") if t != "sel"]
            if iid in selected:
                tags.append("sel")
            self.item(iid, tags=tags)


class FinanceApp:
    NAV_ITEMS = ("工作台", "账务", "账簿", "资金", "往来", "费用票据", "邮件收票", "进销存", "报表", "期末处理", "设置")

    def __init__(self, root: tk.Tk):
        self.root = root
        db.init_db()
        self._run_startup_backup()
        self.default_actor = os.environ.get("USERNAME") or os.environ.get("USER") or "本机用户"
        self._dirty = False
        self._dirty_label = "当前表单"
        self._current_save = None
        self._search_widget = None
        self._active_nav = ""
        self._issues: list[dict] = []
        self._nav_buttons: dict[str, tk.Button] = {}
        self._report_action_buttons: list[tk.Widget] = []
        self._report_cache: dict | None = None
        self._invoice_summary_path: str | None = None
        self._pending_ledger_account: tuple[str, str] | None = None
        self._tray_icon = None
        self._tray_thread = None
        self._close_to_tray = True

        root.title("财务系统 · 可靠单机记账版")
        screen_w, screen_h = root.winfo_screenwidth(), root.winfo_screenheight()
        width, height = min(1220, max(900, screen_w - 48)), min(760, max(590, screen_h - 80))
        root.geometry(f"{width}x{height}+{max(0, (screen_w-width)//2)}+{max(0, (screen_h-height)//3)}")
        root.minsize(min(1180, width), min(700, height))
        root.configure(bg=C_BG)
        root.protocol("WM_DELETE_WINDOW", self._on_close)
        self._start_tray()
        root.bind_all("<Control-s>", self._keyboard_save, add="+")
        root.bind_all("<Control-f>", self._keyboard_search, add="+")
        root.bind_all("<Escape>", self._keyboard_escape, add="+")

        icon = ROOT / "icon.ico"
        if icon.exists():
            try:
                root.iconbitmap(str(icon))
            except tk.TclError:
                pass

        configure_dark_style(root)
        root.grid_rowconfigure(0, weight=1)
        root.grid_columnconfigure(1, weight=1)

        # Sidebar
        self.nav = tk.Frame(root, bg=C_NAV, width=200)
        self.nav.grid(row=0, column=0, sticky="nsew")
        self.nav.grid_propagate(False)
        self.main = tk.Frame(root, bg=C_BG)
        self.main.grid(row=0, column=1, sticky="nsew")
        self.main.grid_columnconfigure(0, weight=1)
        self.main.grid_rowconfigure(2, weight=1)

        self.company_var = tk.StringVar(value="未设置公司")
        self._load_company_name()

        # Sidebar header
        hdr = tk.Frame(self.nav, bg=C_NAV, height=80)
        hdr.pack(fill="x", padx=0, pady=(0, 8))
        hdr.pack_propagate(False)
        tk.Label(
            hdr, text="财务系统", font=FONT_H2,
            bg=C_NAV, fg=C_TEXT,
        ).pack(anchor="w", padx=20, pady=(22, 0))
        tk.Label(
            hdr, textvariable=self.company_var, font=FONT_SMALL,
            bg=C_NAV, fg=C_TEXT_MUTED, wraplength=160, justify="left",
        ).pack(anchor="w", padx=20, pady=(2, 0))

        # Nav buttons
        for name in self.NAV_ITEMS:
            btn = NavButton(self.nav, text=name, command=lambda key=name: self._show_page(key))
            btn.pack(fill="x", padx=8, pady=1)
            self._nav_buttons[name] = btn

        # Sidebar footer
        tk.Label(
            self.nav, text="本机数据 · SQLite\nCtrl+S 保存  Ctrl+F 查找",
            font=("Microsoft YaHei UI", 8), bg=C_NAV, fg=C_TEXT_DIM, justify="left",
        ).pack(side="bottom", anchor="w", padx=20, pady=18)

        # Header bar (dark, no border)
        self.header = tk.Frame(self.main, bg=C_CARD, height=64)
        self.header.grid(row=0, column=0, sticky="ew")
        self.header.grid_propagate(False)
        self.page_title = tk.Label(self.header, text="", font=FONT_TITLE, bg=C_CARD, fg=C_TEXT)
        self.page_title.pack(side="left", padx=24, pady=14)
        self.header_status = tk.Label(self.header, text="", font=FONT_SMALL, bg=C_CARD, fg=C_TEXT_MUTED)
        self.header_status.pack(side="right", padx=24)

        # Issue warning bar (dark amber)
        self.issue_bar = tk.Frame(self.main, bg=C_WARN_SOFT, highlightthickness=0)
        self.issue_text = tk.Label(
            self.issue_bar, text="", bg=C_WARN_SOFT, fg=C_WARN, font=FONT_B,
            anchor="w", justify="left",
        )
        self.issue_text.pack(side="left", fill="x", expand=True, padx=16, pady=10)
        ModernButton(
            self.issue_bar, text="立即处理", variant="danger",
            font=FONT_SMALL, height=28,
            command=self._open_issue_dialog,
        ).pack(side="right", padx=12, pady=8)
        self.issue_bar.grid(row=1, column=0, sticky="ew")

        self.content = tk.Frame(self.main, bg=C_BG, padx=20, pady=16)
        self.content.grid(row=2, column=0, sticky="nsew")
        self._refresh_issue_state()
        self._show_page("工作台", force=True)

    # ------------------------------------------------------------------
    # TD-4: 启动时自动备份
    # ------------------------------------------------------------------
    def _run_startup_backup(self) -> None:
        """Check if today's backup exists; if not, create one and keep N most recent."""
        try:
            from db import backup_database, get_conn as _db_get_conn
            from datetime import date

            # Determine the backup directory from app_paths or fallback to DB parent.
            try:
                import app_paths as _ap
                backup_dir = _ap.BACKUP_DIR
            except Exception:
                backup_dir = Path(db.DB_PATH).resolve().parent.parent / "backups"

            backup_dir.mkdir(parents=True, exist_ok=True)

            today_str = date.today().isoformat()  # YYYY-MM-DD
            existing_today = any(
                f.name.startswith(f"finance-{today_str}")
                for f in backup_dir.iterdir()
                if f.is_file() and f.suffix == ".db"
            )

            if not existing_today:
                backup_database()

            # Keep the most recent N backups, delete older ones.
            MAX_BACKUPS = 5
            backups = sorted(
                [f for f in backup_dir.iterdir() if f.is_file() and f.suffix == ".db"],
                key=lambda f: f.stat().st_mtime,
                reverse=True,
            )
            for old_backup in backups[MAX_BACKUPS:]:
                old_backup.unlink(missing_ok=True)
        except Exception:
            # Never let a backup failure prevent the app from starting.
            pass

    def _configure_style(self) -> None:
        configure_dark_style(self.root)

    def _load_company_name(self) -> None:
        try:
            conn = db.get_conn()
            row = conn.execute("SELECT name FROM company WHERE id=1").fetchone()
            conn.close()
            if row and row["name"]:
                self.company_var.set(row["name"])
        except Exception:
            pass

    def _keyboard_save(self, _event=None):
        if callable(self._current_save):
            self._current_save()
            return "break"
        return None

    def _keyboard_search(self, _event=None):
        if self._search_widget and self._search_widget.winfo_exists():
            self._search_widget.focus_set()
            try:
                self._search_widget.selection_range(0, tk.END)
            except tk.TclError:
                pass
            return "break"
        return None

    def _keyboard_escape(self, _event=None):
        focused = self.root.focus_get()
        if isinstance(focused, (tk.Entry, ttk.Entry)):
            self.root.focus_set()
        return None

    def _on_close(self) -> None:
        if not self._close_to_tray:
            if self._confirm_abandon():
                self._quit_app()
            return
        # 最小化到托盘
        self.root.withdraw()
        if self._tray_icon:
            try:
                self._tray_icon.visible = True
            except Exception:
                pass

    def _show_window(self, icon=None, item=None) -> None:
        self.root.deiconify()
        self.root.lift()
        self.root.focus_force()
        if self._tray_icon:
            try:
                self._tray_icon.visible = False
            except Exception:
                pass

    def _hide_window(self, icon=None, item=None) -> None:
        self.root.withdraw()
        if self._tray_icon:
            try:
                self._tray_icon.visible = True
            except Exception:
                pass

    def _quit_app(self, icon=None, item=None) -> None:
        self._close_to_tray = False
        self._stop_tray()
        if self.root.winfo_exists():
            self.root.destroy()

    def _start_tray(self) -> None:
        try:
            from PIL import Image
            import pystray
            import threading

            icon_path = ROOT / "icon.ico"
            if not icon_path.exists():
                return
            image = Image.open(str(icon_path))
            menu = pystray.Menu(
                pystray.MenuItem("显示窗口", self._show_window),
                pystray.MenuItem("隐藏窗口", self._hide_window),
                pystray.MenuItem("退出", self._quit_app),
            )
            icon = pystray.Icon("finance", image, "财务系统", menu)
            icon.visible = False
            self._tray_icon = icon

            def _run():
                icon.run()

            t = threading.Thread(target=_run, daemon=True)
            t.start()
            self._tray_thread = t
        except Exception:
            self._tray_icon = None
            self._tray_thread = None

    def _stop_tray(self) -> None:
        if self._tray_icon:
            try:
                self._tray_icon.stop()
            except Exception:
                pass
            self._tray_icon = None
        self._tray_thread = None

    def _mark_dirty(self, label: str = "当前表单") -> None:
        self._dirty = True
        self._dirty_label = label
        self.header_status.config(text="有未保存更改", fg=C_WARN)

    def _clear_dirty(self) -> None:
        self._dirty = False
        self._dirty_label = "当前表单"
        self.header_status.config(text="数据保存在本机", fg=C_TEXT_MUTED)

    def _confirm_abandon(self) -> bool:
        if not self._dirty:
            return True
        keep_going = messagebox.askyesno(
            "未保存更改", f"{self._dirty_label}尚未保存。是否放弃更改？", parent=self.root
        )
        if keep_going:
            self._clear_dirty()
        return keep_going

    def _show_page(self, name: str, *, force: bool = False) -> None:
        if not force and name != self._active_nav and not self._confirm_abandon():
            return
        self._active_nav = name
        self._current_save = None
        self._search_widget = None
        for key, button in self._nav_buttons.items():
            button.set_active(key == name)
        for child in self.content.winfo_children():
            child.destroy()
        self.page_title.configure(text=name)
        self._refresh_issue_state()
        builders = {
            "工作台": self._page_dashboard,
            "账务": self._page_ledger,
            "账簿": self._page_books,
            "资金": self._page_cash,
            "往来": self._page_partners,
            "费用票据": self._page_expense_invoice,
            "邮件收票": self._page_mail,
            "进销存": self._page_inventory,
            "报表": self._page_reports,
            "期末处理": self._page_closing,
            "设置": self._page_settings,
        }
        try:
            builders[name](self.content)
        except Exception as exc:
            self._show_page_error(exc)

    def _show_page_error(self, exc: Exception) -> None:
        panel = tk.Frame(self.content, bg=C_PANEL, padx=24, pady=24)
        panel.pack(fill="x")
        tk.Label(panel, text="页面载入失败", font=FONT_B, bg=C_PANEL, fg=C_DANGER).pack(anchor="w")
        tk.Label(
            panel, text=str(exc), font=FONT, bg=C_PANEL, fg=C_TEXT,
            justify="left", wraplength=760,
        ).pack(anchor="w", pady=(8, 0))

    def _open_issues(self) -> list[dict]:
        try:
            getter = getattr(db, "get_open_data_issues", None)
            if getter:
                return [dict(row) for row in getter(severity="blocking")]
            return [dict(row) for row in reports.get_blocking_issues()]
        except Exception:
            return []

    def _refresh_issue_state(self) -> None:
        self._issues = self._open_issues()
        if self._issues:
            self.issue_text.configure(text=f"发现 {len(self._issues)} 项阻断性数据异常；正式报表已暂停。")
            self.issue_bar.grid()
        else:
            self.issue_bar.grid_remove()
        for button in self._report_action_buttons:
            if button.winfo_exists():
                button.configure(state="disabled" if self._issues else "normal")

    def _panel(self, parent, **kwargs) -> tk.Frame:
        options = {"bg": C_CARD, "highlightthickness": 1, "highlightbackground": C_BORDER}
        options.update(kwargs)
        return tk.Frame(parent, **options)

    def _tree(self, parent, columns, *, height=12, selectmode="browse") -> ttk.Treeview:
        wrapper = tk.Frame(parent, bg=C_CARD)
        wrapper.pack(fill="both", expand=True)
        tree = _ZebraTree(wrapper, columns=[item[0] for item in columns], show="headings", height=height, selectmode=selectmode)
        ybar = ttk.Scrollbar(wrapper, orient="vertical", command=tree.yview)
        xbar = ttk.Scrollbar(wrapper, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=ybar.set, xscrollcommand=xbar.set)
        tree.grid(row=0, column=0, sticky="nsew")
        ybar.grid(row=0, column=1, sticky="ns")
        xbar.grid(row=1, column=0, sticky="ew")
        wrapper.grid_rowconfigure(0, weight=1)
        wrapper.grid_columnconfigure(0, weight=1)
        for key, label, width, anchor in columns:
            tree.heading(key, text=label)
            tree.column(key, width=width, minwidth=width, anchor=anchor, stretch=True)
        return tree

    def _period_values(self) -> list[str]:
        values = {_current_period()}
        try:
            conn = db.get_conn()
            if _table_exists(conn, "voucher"):
                values.update(row[0] for row in conn.execute("SELECT DISTINCT period FROM voucher WHERE period IS NOT NULL"))
            if _table_exists(conn, "period"):
                values.update(f"{int(row[0]):04d}-{int(row[1]):02d}" for row in conn.execute("SELECT year,month FROM period"))
            conn.close()
        except Exception:
            pass
        return sorted((str(v) for v in values if v), reverse=True)

    def _accounts(self, *, category: str | None = None, direction: str | None = None) -> list[dict]:
        conn = db.get_conn()
        sql, args = "SELECT * FROM account WHERE COALESCE(is_leaf,1)=1", []
        if category:
            sql += " AND category=?"
            args.append(category)
        if direction:
            sql += " AND direction=?"
            args.append(direction)
        sql += " ORDER BY code"
        rows = [dict(row) for row in conn.execute(sql, args).fetchall()]
        conn.close()
        return rows

    @staticmethod
    def _account_display(account: dict) -> str:
        return f"{account['code']}  {account['name']}"

    # ------------------------------ 工作台 ------------------------------
    def _page_dashboard(self, parent) -> None:
        self.header_status.configure(text=f"今天 {_today()} · 数据保存在本机", fg=C_TEXT_MUTED)

        # Top heading row
        heading = tk.Frame(parent, bg=C_BG)
        heading.pack(fill="x", pady=(0, 16))
        tk.Label(heading, text="今日概览", font=FONT_H2, bg=C_BG, fg=C_TEXT).pack(side="left")
        ModernButton(heading, text="⟳ 刷新", variant="ghost",
                     font=FONT_SMALL, height=28,
                     command=lambda: self._show_page("工作台", force=True)).pack(side="right")

        # KPI stat cards
        counts = {"待审核凭证": 0, "未处理报销": 0, "待确认发票": 0}
        try:
            conn = db.get_conn()
            if _table_exists(conn, "voucher"):
                counts["待审核凭证"] = conn.execute("SELECT COUNT(*) FROM voucher WHERE status IN ('待审核','submitted','pending_review')").fetchone()[0]
            if _table_exists(conn, "expense"):
                counts["未处理报销"] = conn.execute("SELECT COUNT(*) FROM expense WHERE status NOT IN ('已付款','paid','已驳回','rejected')").fetchone()[0]
            if _table_exists(conn, "invoice_doc"):
                counts["待确认发票"] = conn.execute("SELECT COUNT(*) FROM invoice_doc WHERE status IN ('待确认','unconfirmed')").fetchone()[0]
            conn.close()
        except Exception:
            pass

        cards = tk.Frame(parent, bg=C_BG)
        cards.pack(fill="x")
        items = [
            ("阻断异常", str(len(self._issues)), C_DANGER if self._issues else C_OK),
            *[(k, str(v), C_ACCENT) for k, v in counts.items()],
        ]
        for index, (label, value, colour) in enumerate(items):
            sc = StatCard(cards, label=label, value=value, colour=colour)
            sc.grid(row=0, column=index, sticky="ew", padx=(0 if index == 0 else 10, 0))
            cards.grid_columnconfigure(index, weight=1)

        # Lower section: recent vouchers + quick actions
        lower = tk.Frame(parent, bg=C_BG)
        lower.pack(fill="both", expand=True, pady=(18, 0))
        lower.grid_columnconfigure(0, weight=3)
        lower.grid_columnconfigure(1, weight=1)

        # Recent vouchers panel
        recent_panel = CardFrame(lower, radius=12, bg=C_CARD, border=C_BORDER, padx=14, pady=14)
        recent_panel.grid(row=0, column=0, sticky="nsew", padx=(0, 10))
        tk.Label(recent_panel.inner, text="最近凭证", font=FONT_H3, bg=C_CARD, fg=C_TEXT).pack(anchor="w", pady=(0, 10))
        tree = self._tree(recent_panel.inner, (("vno", "凭证号", 140, "w"), ("date", "日期", 90, "center"), ("summary", "摘要", 220, "w"), ("status", "状态", 70, "center")), height=10)
        try:
            for record in self._voucher_records(limit=12):
                head = record["voucher"]
                tree.insert("", "end", values=(head.get("vno"), head.get("date"), head.get("summary"), _status(head.get("status"))))
        except Exception:
            pass

        # Quick actions panel
        quick = CardFrame(lower, radius=12, bg=C_CARD, border=C_BORDER, padx=16, pady=16)
        quick.grid(row=0, column=1, sticky="nsew")
        tk.Label(quick.inner, text="快捷操作", font=FONT_H3, bg=C_CARD, fg=C_TEXT).pack(anchor="w", pady=(0, 12))
        for label, page in (("录入凭证", "账务"), ("新增收付款", "资金"), ("核对报表", "报表"), ("备份与设置", "设置")):
            ModernButton(quick.inner, text=label, variant="ghost",
                         font=FONT_SMALL, height=32,
                         command=lambda p=page: self._show_page(p)).pack(fill="x", pady=4)
        if self._issues:
            ModernButton(quick.inner, text="处理数据异常", variant="danger",
                         font=FONT_SMALL, height=32,
                         command=self._open_issue_dialog).pack(fill="x", pady=(16, 4))

    # ------------------------------ 账务 ------------------------------
    def _page_ledger(self, parent) -> None:
        self.header_status.configure(text="凭证按 草稿 → 待审核 → 已审核 → 已记账 流转", fg=C_TEXT_MUTED)
        notebook = ttk.Notebook(parent)
        notebook.pack(fill="both", expand=True)
        tabs = [tk.Frame(notebook, bg=C_PANEL) for _ in range(4)]
        for tab, label in zip(tabs, ("凭证录入", "凭证查询", "科目余额", "科目明细")):
            notebook.add(tab, text=label)
        self._build_voucher_editor(tabs[0])
        self._build_voucher_browser(tabs[1])
        self._build_trial_balance(tabs[2])
        self._build_account_ledger(tabs[3])
        previous = {"index": 0, "guard": False}

        def tab_changed(_event=None):
            if previous["guard"]:
                return
            selected = notebook.index(notebook.select())
            if previous["index"] == 0 and selected != 0 and self._dirty:
                if not self._confirm_abandon():
                    previous["guard"] = True
                    notebook.select(0)
                    self.root.after_idle(lambda: previous.update(guard=False))
                    return
            previous["index"] = selected

        notebook.bind("<<NotebookTabChanged>>", tab_changed)
        if self._pending_ledger_account:
            account_code, period = self._pending_ledger_account
            self._pending_ledger_account = None
            if hasattr(self, "ledger_account_var"):
                match = next((value for value in self.ledger_account_combo["values"] if str(value).startswith(account_code + " ")), None)
                if match:
                    self.ledger_account_var.set(match)
                self.ledger_period_var.set(period)
                notebook.select(3)
                self.root.after_idle(self._refresh_account_ledger)

    def _build_voucher_editor(self, parent) -> None:
        lines: list[dict] = []
        top = tk.Frame(parent, bg=C_PANEL, padx=14, pady=12)
        top.pack(fill="x")
        date_var = tk.StringVar(value=_today())
        summary_var = tk.StringVar()
        actor_var = tk.StringVar(value=self.default_actor)
        fields = (("日期", date_var, 13), ("摘要", summary_var, 34), ("制单人", actor_var, 16))
        for index, (label, var, width) in enumerate(fields):
            tk.Label(top, text=label, font=FONT, bg=C_PANEL).grid(row=0, column=index * 2, padx=(0, 5), pady=2, sticky="e")
            ttk.Entry(top, textvariable=var, width=width).grid(row=0, column=index * 2 + 1, padx=(0, 14), pady=2, sticky="ew")
        top.grid_columnconfigure(3, weight=1)
        total_var = tk.StringVar(value="借 0.00  /  贷 0.00")
        tk.Label(top, textvariable=total_var, font=FONT_B, bg=C_PANEL, fg=C_ACCENT).grid(row=1, column=0, columnspan=6, sticky="e", pady=(8, 0))

        table_panel = tk.Frame(parent, bg=C_PANEL, padx=14, pady=8)
        table_panel.pack(fill="both", expand=True)
        tree = self._tree(table_panel, (("account", "科目", 190, "w"), ("summary", "摘要", 240, "w"), ("debit", "借方", 105, "e"), ("credit", "贷方", 105, "e"), ("aux", "辅助核算", 120, "w")), height=13)

        def refresh_lines():
            tree.delete(*tree.get_children())
            for index, line in enumerate(lines):
                account = next((a for a in self._accounts() if a["code"] == line["account_code"]), None)
                label = self._account_display(account) if account else line["account_code"]
                aux = line.get("aux_type") or "无"
                if line.get("aux_id"):
                    aux += f" #{line['aux_id']}"
                tree.insert("", "end", iid=str(index), values=(label, line.get("summary", ""), _money(line.get("debit_cents")), _money(line.get("credit_cents")), aux))
            debit = sum(int(line.get("debit_cents") or 0) for line in lines)
            credit = sum(int(line.get("credit_cents") or 0) for line in lines)
            mark = " ✓" if debit and debit == credit else ""
            total_var.set(f"借 {_money(debit)}  /  贷 {_money(credit)}{mark}")

        def edit_line(index: int | None = None):
            initial = lines[index] if index is not None else None
            dialog = LineDialog(self.root, self._accounts(), default_summary=summary_var.get(), initial=initial)
            if dialog.result:
                if index is None:
                    lines.append(dialog.result)
                else:
                    lines[index] = dialog.result
                refresh_lines()
                self._mark_dirty("凭证")

        def selected_index() -> int | None:
            selection = tree.selection()
            return int(selection[0]) if selection else None

        def delete_line():
            index = selected_index()
            if index is None:
                messagebox.showwarning("请选择分录", "请先选择要删除的分录。", parent=self.root)
                return
            del lines[index]
            refresh_lines()
            self._mark_dirty("凭证")

        tree.bind("<Double-1>", lambda _e: edit_line(selected_index()) if selected_index() is not None else None)
        actions = tk.Frame(parent, bg=C_PANEL, padx=14, pady=10)
        actions.pack(fill="x")
        ttk.Button(actions, text="增加分录", command=lambda: edit_line()).pack(side="left", padx=(0, 6))
        ttk.Button(actions, text="编辑分录", command=lambda: edit_line(selected_index()) if selected_index() is not None else None).pack(side="left", padx=6)
        ttk.Button(actions, text="删除分录", command=delete_line).pack(side="left", padx=6)

        idempotency = {"key": uuid.uuid4().hex}

        def save(submit: bool):
            try:
                date_value = _validate_date(date_var.get())
                if not summary_var.get().strip():
                    raise ValueError("凭证摘要不能为空")
                if not actor_var.get().strip():
                    raise ValueError("制单人不能为空")
                debit = sum(int(line.get("debit_cents") or 0) for line in lines)
                credit = sum(int(line.get("credit_cents") or 0) for line in lines)
                if len(lines) < 2:
                    raise ValueError("凭证至少需要两行分录")
                if debit <= 0 or debit != credit:
                    raise ValueError("借贷必须相等且金额不能为零")
                if getattr(ledger, "is_period_closed", lambda _p: False)(date_value[:7]):
                    raise ValueError("该会计期间已结账，不能新增凭证")
                result = ledger.create_voucher(
                    date_value,
                    summary_var.get().strip(),
                    list(lines),
                    creator=actor_var.get().strip(),
                    source="manual",
                    idempotency_key=idempotency["key"],
                    submit=submit,
                )
                ok, message, voucher_id = _result(result)
                if not ok:
                    raise ValueError(message)
            except Exception as exc:
                messagebox.showerror("保存失败", str(exc), parent=self.root)
                return
            messagebox.showinfo("凭证已保存", message or f"凭证 #{voucher_id} 已保存。", parent=self.root)
            lines.clear()
            summary_var.set("")
            idempotency["key"] = uuid.uuid4().hex
            refresh_lines()
            self._clear_dirty()

        ttk.Button(actions, text="保存草稿", command=lambda: save(False)).pack(side="right", padx=(6, 0))
        tk.Button(actions, text="保存并提交审核", font=FONT_B, bg=C_ACCENT, fg="white", relief="flat", padx=14, pady=7, command=lambda: save(True)).pack(side="right", padx=6)
        self._current_save = lambda: save(False)
        for variable in (date_var, summary_var, actor_var):
            variable.trace_add("write", lambda *_args: self._mark_dirty("凭证"))

    def _voucher_records(self, period=None, status=None, limit=200) -> list[dict]:
        getter = getattr(ledger, "get_vouchers", None)
        if getter:
            try:
                rows = _call_supported(getter, period=period, status=status, limit=limit)
            except TypeError:
                rows = getter(period, status, limit)
            records = []
            for row in rows:
                if isinstance(row, dict) and "voucher" in row:
                    records.append({"voucher": dict(row["voucher"]), "lines": [dict(x) for x in row.get("lines", [])]})
                else:
                    records.append({"voucher": dict(row), "lines": []})
            return records
        return [{"voucher": dict(row), "lines": []} for row in ledger.list_vouchers(period)]

    def _build_voucher_browser(self, parent) -> None:
        filters = tk.Frame(parent, bg=C_PANEL, padx=12, pady=10)
        filters.pack(fill="x")
        period_var = tk.StringVar(value=_current_period())
        status_var = tk.StringVar(value="全部")
        search_var = tk.StringVar()
        tk.Label(filters, text="期间", bg=C_PANEL, font=FONT).pack(side="left")
        ttk.Combobox(filters, textvariable=period_var, values=["全部", *self._period_values()], width=11, state="readonly").pack(side="left", padx=(5, 12))
        tk.Label(filters, text="状态", bg=C_PANEL, font=FONT).pack(side="left")
        ttk.Combobox(filters, textvariable=status_var, values=("全部", "草稿", "待审核", "已审核", "已记账", "已冲销"), width=10, state="readonly").pack(side="left", padx=(5, 12))
        tk.Label(filters, text="查找", bg=C_PANEL, font=FONT).pack(side="left")
        search = ttk.Entry(filters, textvariable=search_var, width=25)
        search.pack(side="left", padx=(5, 8), fill="x", expand=True)
        self._search_widget = search
        ttk.Button(filters, text="查询", command=lambda: refresh()).pack(side="left")

        panes = ttk.Panedwindow(parent, orient="vertical")
        panes.pack(fill="both", expand=True, padx=12, pady=(0, 8))
        upper, lower = tk.Frame(panes, bg=C_PANEL), tk.Frame(panes, bg=C_PANEL)
        panes.add(upper, weight=3)
        panes.add(lower, weight=2)
        tree = self._tree(upper, (("id", "ID", 55, "center"), ("vno", "凭证号", 145, "w"), ("date", "日期", 95, "center"), ("summary", "摘要", 250, "w"), ("creator", "制单人", 90, "w"), ("source", "来源", 85, "center"), ("status", "状态", 80, "center")), height=9)
        detail = self._tree(lower, (("seq", "序号", 50, "center"), ("account", "科目", 190, "w"), ("summary", "分录摘要", 240, "w"), ("debit", "借方", 105, "e"), ("credit", "贷方", 105, "e"), ("aux", "辅助核算", 100, "w")), height=6)
        state: dict[str, object] = {"records": {}, "selected": None}

        def refresh():
            tree.delete(*tree.get_children())
            detail.delete(*detail.get_children())
            state["records"] = {}
            period = None if period_var.get() in ("", "全部") else _validate_period(period_var.get())
            selected_status = None if status_var.get() == "全部" else VOUCHER_STATUS_CODES.get(status_var.get(), status_var.get())
            keyword = search_var.get().strip().lower()
            try:
                records = self._voucher_records(period=period, status=selected_status, limit=500)
            except Exception as exc:
                messagebox.showerror("查询失败", str(exc), parent=self.root)
                return
            for record in records:
                head = record["voucher"]
                haystack = " ".join(str(head.get(key) or "") for key in ("vno", "summary", "creator", "source")).lower()
                if keyword and keyword not in haystack:
                    continue
                row_id = str(head.get("id"))
                state["records"][row_id] = record
                tree.insert("", "end", iid=row_id, values=(head.get("id"), head.get("vno"), head.get("date"), head.get("summary"), head.get("creator"), head.get("source"), _status(head.get("status"))))
            update_buttons()

        def selected_record():
            selection = tree.selection()
            return state["records"].get(selection[0]) if selection else None

        def show_detail(_event=None):
            detail.delete(*detail.get_children())
            record = selected_record()
            if not record:
                state["selected"] = None
                update_buttons()
                return
            state["selected"] = record
            for index, line in enumerate(record.get("lines", []), 1):
                debit = _row_cents(line, "debit_cents", "debit")
                credit = _row_cents(line, "credit_cents", "credit")
                account = f"{line.get('account_code') or ''}  {line.get('account_name') or ''}".strip()
                aux = line.get("aux_type") or "无"
                if line.get("aux_id"):
                    aux += f" #{line['aux_id']}"
                detail.insert("", "end", values=(line.get("seq", index), account, line.get("summary", ""), _money(debit), _money(credit), aux))
            update_buttons()

        tree.bind("<<TreeviewSelect>>", show_detail)
        search.bind("<Return>", lambda _e: refresh())
        period_var.trace_add("write", lambda *_a: self.root.after_idle(refresh))
        status_var.trace_add("write", lambda *_a: self.root.after_idle(refresh))

        buttons = tk.Frame(parent, bg=C_PANEL, padx=12, pady=9)
        buttons.pack(fill="x")
        submit_btn = ttk.Button(buttons, text="提交审核")
        audit_btn = ttk.Button(buttons, text="审核通过")
        post_btn = ttk.Button(buttons, text="记账")
        reverse_btn = ttk.Button(buttons, text="生成冲销凭证")
        delete_btn = ttk.Button(buttons, text="删除草稿")
        for button in (delete_btn, reverse_btn, post_btn, audit_btn, submit_btn):
            button.pack(side="right", padx=4)

        def update_buttons():
            record = selected_record()
            status = _status(record["voucher"].get("status")) if record else ""
            submit_btn.configure(state="normal" if status == "草稿" else "disabled")
            audit_btn.configure(state="normal" if status == "待审核" else "disabled")
            post_btn.configure(state="normal" if status == "已审核" else "disabled")
            reverse_btn.configure(state="normal" if status == "已记账" else "disabled")
            delete_btn.configure(state="normal" if status == "草稿" else "disabled")

        def workflow(action: str):
            record = selected_record()
            if not record:
                messagebox.showwarning("请选择凭证", "请先选择一张凭证。", parent=self.root)
                return
            head = record["voucher"]
            dialog = WorkflowDialog(self.root, action, self.default_actor)
            if not dialog.result:
                return
            if action == "删除草稿" and not messagebox.askyesno(
                "确认删除草稿",
                f"确定删除草稿 {head.get('vno', '')}？此操作只删除未提交审核的草稿，不能恢复。",
                parent=self.root,
            ):
                return
            try:
                if action == "提交审核":
                    result = _call_supported(ledger.submit_voucher, voucher_id=head["id"], actor=dialog.result["actor"])
                elif action == "审核通过":
                    result = _call_supported(ledger.audit_voucher, voucher_id=head["id"], reviewer=dialog.result["actor"], actor=dialog.result["actor"], exception_reason=dialog.result["reason"] or None)
                elif action == "记账":
                    result = _call_supported(ledger.post_voucher, voucher_id=head["id"], poster=dialog.result["actor"], actor=dialog.result["actor"], exception_reason=dialog.result["reason"] or None)
                elif action == "删除草稿":
                    result = _call_supported(ledger.delete_draft_voucher, voucher_id=head["id"], actor=dialog.result["actor"], reason=dialog.result["reason"])
                else:
                    result = _call_supported(ledger.reverse_voucher, voucher_id=head["id"], date_str=dialog.result["date"], actor=dialog.result["actor"], reason=dialog.result["reason"])
                ok, message, _data = _result(result)
                if not ok:
                    raise ValueError(message)
            except Exception as exc:
                messagebox.showerror(f"{action}失败", str(exc), parent=self.root)
                return
            messagebox.showinfo(action, message or "操作成功。", parent=self.root)
            refresh()
            self._refresh_issue_state()

        submit_btn.configure(command=lambda: workflow("提交审核"))
        audit_btn.configure(command=lambda: workflow("审核通过"))
        post_btn.configure(command=lambda: workflow("记账"))
        reverse_btn.configure(command=lambda: workflow("冲销"))
        delete_btn.configure(command=lambda: workflow("删除草稿"))
        refresh()

    def _normalize_trial(self, value) -> dict:
        if isinstance(value, dict):
            rows = [dict(row) for row in value.get("rows", value.get("accounts", []))]
            debit = value.get("total_debit_cents")
            credit = value.get("total_credit_cents")
            if debit is None:
                debit = _parse_cents(value.get("total_debit", 0))
            if credit is None:
                credit = _parse_cents(value.get("total_credit", 0))
            return {"rows": rows, "debit": int(debit), "credit": int(credit), "balanced": bool(value.get("balanced", int(debit) == int(credit)))}
        rows, debit, credit = value
        return {"rows": [dict(row) for row in rows], "debit": _parse_cents(debit), "credit": _parse_cents(credit), "balanced": abs(float(debit) - float(credit)) < 0.005}

    def _build_trial_balance(self, parent) -> None:
        controls = tk.Frame(parent, bg=C_PANEL, padx=12, pady=10)
        controls.pack(fill="x")
        period_var = tk.StringVar(value=_current_period())
        tk.Label(controls, text="截至期间", bg=C_PANEL, font=FONT).pack(side="left")
        ttk.Combobox(controls, textvariable=period_var, values=self._period_values(), width=12, state="readonly").pack(side="left", padx=6)
        summary_var = tk.StringVar()
        tk.Label(controls, textvariable=summary_var, bg=C_PANEL, font=FONT_B, fg=C_TEXT).pack(side="right")
        table_panel = tk.Frame(parent, bg=C_PANEL, padx=12, pady=10)
        table_panel.pack(fill="both", expand=True)
        tree = self._tree(table_panel, (("code", "科目", 90, "w"), ("name", "名称", 230, "w"), ("category", "类别", 80, "center"), ("debit", "借方余额", 130, "e"), ("credit", "贷方余额", 130, "e")), height=16)

        def refresh():
            tree.delete(*tree.get_children())
            try:
                period = _validate_period(period_var.get())
                result = self._normalize_trial(ledger.trial_balance(period))
                for row in result["rows"]:
                    debit = row.get("debit_balance_cents")
                    credit = row.get("credit_balance_cents")
                    if debit is None:
                        debit = _parse_cents(row.get("debit_balance", 0))
                    if credit is None:
                        credit = _parse_cents(row.get("credit_balance", 0))
                    tree.insert("", "end", values=(row.get("code") or row.get("account_code"), row.get("name") or row.get("account_name"), row.get("category", ""), _money(debit), _money(credit)))
                symbol = "平衡 ✓" if result["balanced"] else "不平衡 ✗"
                summary_var.set(f"借 {_money(result['debit'])}  贷 {_money(result['credit'])}  ·  {symbol}")
            except Exception as exc:
                messagebox.showerror("生成失败", str(exc), parent=self.root)

        ttk.Button(controls, text="生成", command=refresh).pack(side="left", padx=4)
        refresh()

    def _build_account_ledger(self, parent) -> None:
        accounts = self._accounts()
        values = [self._account_display(account) for account in accounts]
        controls = tk.Frame(parent, bg=C_PANEL, padx=12, pady=10)
        controls.pack(fill="x")
        self.ledger_account_var = tk.StringVar(value=values[0] if values else "")
        self.ledger_period_var = tk.StringVar(value=_current_period())
        tk.Label(controls, text="科目", bg=C_PANEL, font=FONT).pack(side="left")
        self.ledger_account_combo = ttk.Combobox(controls, textvariable=self.ledger_account_var, values=values, width=27, state="readonly")
        self.ledger_account_combo.pack(side="left", padx=6)
        tk.Label(controls, text="期间", bg=C_PANEL, font=FONT).pack(side="left", padx=(10, 0))
        ttk.Combobox(controls, textvariable=self.ledger_period_var, values=self._period_values(), width=11, state="readonly").pack(side="left", padx=6)
        table_panel = tk.Frame(parent, bg=C_PANEL, padx=12, pady=10)
        table_panel.pack(fill="both", expand=True)
        self.account_ledger_tree = self._tree(table_panel, (("date", "日期", 95, "center"), ("vno", "凭证号", 145, "w"), ("summary", "摘要", 260, "w"), ("debit", "借方", 105, "e"), ("credit", "贷方", 105, "e"), ("balance", "余额", 115, "e")), height=16)
        ttk.Button(controls, text="查询", command=self._refresh_account_ledger).pack(side="left", padx=4)

    def _refresh_account_ledger(self) -> None:
        tree = getattr(self, "account_ledger_tree", None)
        if not tree:
            return
        tree.delete(*tree.get_children())
        try:
            code = self.ledger_account_var.get().split()[0]
            period = _validate_period(self.ledger_period_var.get())
            result = ledger.ledger(code, period)
            for row in result.get("entries", []):
                debit = _row_cents(row, "debit_cents", "debit")
                credit = _row_cents(row, "credit_cents", "credit")
                balance = row.get("balance_cents")
                if balance is None:
                    balance = _parse_cents(abs(float(row.get("balance") or 0))) * (-1 if float(row.get("balance") or 0) < 0 else 1)
                tree.insert("", "end", values=(row.get("date"), row.get("vno"), row.get("summary"), _money(debit), _money(credit), _money(balance)))
        except Exception as exc:
            messagebox.showerror("查询失败", str(exc), parent=self.root)

    # ------------------------------ 资金 ------------------------------
    def _page_cash(self, parent) -> None:
        self.header_status.configure(text="收付款单提交后生成待审核凭证，不会直接记账", fg=C_TEXT_MUTED)
        form = self._panel(parent, padx=16, pady=14)
        form.pack(fill="x", pady=(0, 12))
        doc_type = tk.StringVar(value="收款")
        date_var = tk.StringVar(value=_today())
        amount_var = tk.StringVar()
        actor_var = tk.StringVar(value=self.default_actor)
        summary_var = tk.StringVar()
        accounts = self._accounts()
        account_values = [self._account_display(a) for a in accounts]
        counterpart_var = tk.StringVar(value=next((v for v in account_values if v.startswith("1122 ")), account_values[0] if account_values else ""))
        conn = db.get_conn()
        banks = [dict(row) for row in conn.execute("SELECT * FROM bank_account ORDER BY id").fetchall()] if _table_exists(conn, "bank_account") else []
        conn.close()
        bank_values = [f"{row['id']}|{row.get('name') or row.get('bank') or '银行账户'}|{row.get('account_code') or '1002'}" for row in banks]
        if not bank_values:
            bank_values = ["0|默认银行存款|1002"]
        bank_var = tk.StringVar(value=bank_values[0])
        partners = arap.list_partners()
        partner_values = [f"{row['id']}|{row['type']}|{row['name']}" for row in partners]
        partner_var = tk.StringVar(value=partner_values[0] if partner_values else "")

        fields = [
            ("单据类型", ttk.Combobox(form, textvariable=doc_type, values=("收款", "付款"), width=12, state="readonly")),
            ("业务日期", ttk.Entry(form, textvariable=date_var, width=15)),
            ("金额（元）", ttk.Entry(form, textvariable=amount_var, width=16)),
            ("经办人", ttk.Entry(form, textvariable=actor_var, width=18)),
            ("银行账户", ttk.Combobox(form, textvariable=bank_var, values=bank_values, width=30, state="readonly")),
            ("往来方", ttk.Combobox(form, textvariable=partner_var, values=partner_values, width=30, state="readonly")),
            ("对方科目", ttk.Combobox(form, textvariable=counterpart_var, values=account_values, width=30, state="readonly")),
            ("摘要", ttk.Entry(form, textvariable=summary_var, width=42)),
        ]
        for index, (label, widget) in enumerate(fields):
            row, col = divmod(index, 4)
            tk.Label(form, text=label, font=FONT, bg=C_PANEL).grid(row=row, column=col * 2, padx=(0, 5), pady=6, sticky="e")
            widget.grid(row=row, column=col * 2 + 1, padx=(0, 16), pady=6, sticky="ew")
            form.grid_columnconfigure(col * 2 + 1, weight=1)

        def type_changed(*_args):
            wanted = "1122" if doc_type.get() == "收款" else "2202"
            match = next((value for value in account_values if value.startswith(wanted + " ")), None)
            if match:
                counterpart_var.set(match)

        doc_type.trace_add("write", type_changed)

        history_panel = self._panel(parent, padx=12, pady=10)
        history_panel.pack(fill="both", expand=True)
        tk.Label(history_panel, text="最近收付款单", font=FONT_B, bg=C_PANEL, fg=C_TEXT).pack(anchor="w", pady=(0, 8))
        history = self._tree(history_panel, (("type", "类型", 70, "center"), ("date", "日期", 95, "center"), ("partner", "往来方", 170, "w"), ("summary", "摘要", 260, "w"), ("amount", "金额", 110, "e"), ("status", "状态", 90, "center"), ("voucher", "凭证ID", 80, "center")), height=10)

        def refresh_history():
            history.delete(*history.get_children())
            try:
                conn = db.get_conn()
                columns = _table_columns(conn, "cash_doc")
                if columns:
                    rows = conn.execute("SELECT c.*,p.name partner_name FROM cash_doc c LEFT JOIN partner p ON p.id=c.partner_id ORDER BY c.date DESC,c.id DESC LIMIT 200").fetchall()
                    for row in rows:
                        item = dict(row)
                        amount = _row_cents(item, "amount_cents", "amount")
                        history.insert("", "end", values=(item.get("doc_type"), item.get("date"), item.get("partner_name") or "", item.get("summary"), _money(amount), _status(item.get("status")), item.get("voucher_id") or ""))
                conn.close()
            except Exception as exc:
                messagebox.showerror("载入失败", str(exc), parent=self.root)

        def submit():
            try:
                date_value = _validate_date(date_var.get())
                amount_cents = _parse_cents(amount_var.get(), positive=True)
                if not partner_var.get():
                    raise ValueError("请先在往来管理中新增并选择往来方")
                if not summary_var.get().strip():
                    raise ValueError("摘要不能为空")
                if not actor_var.get().strip():
                    raise ValueError("经办人不能为空")
                if getattr(ledger, "is_period_closed", lambda _p: False)(date_value[:7]):
                    raise ValueError("该会计期间已结账")
                partner_id, partner_type, _name = partner_var.get().split("|", 2)
                bank_id, _bank_name, fund_code = bank_var.get().split("|", 2)
                counterpart_code = counterpart_var.get().split()[0]
                function = ledger.create_cash_doc
                signature = inspect.signature(function)
                kwargs = {
                    "doc_type": doc_type.get(), "date_str": date_value,
                    "partner_type": partner_type, "partner_id": int(partner_id),
                    "amount_cents": amount_cents, "amount": amount_cents / 100.0,
                    "counterparty_account_code": counterpart_code,
                    "account_code": fund_code, "fund_account_code": fund_code,
                    "bank_account_id": None if bank_id == "0" else int(bank_id),
                    "arap_item_id": None, "summary": summary_var.get().strip(),
                    "creator": actor_var.get().strip(), "submit": True,
                    "idempotency_key": f"cash:{uuid.uuid4().hex}",
                }
                result = function(**{key: value for key, value in kwargs.items() if key in signature.parameters or any(p.kind == inspect.Parameter.VAR_KEYWORD for p in signature.parameters.values())})
                ok, message, _data = _result(result)
                if not ok:
                    raise ValueError(message)
            except Exception as exc:
                messagebox.showerror("提交失败", str(exc), parent=self.root)
                return
            messagebox.showinfo("收付款单已提交", message or "已生成待审核凭证。", parent=self.root)
            amount_var.set("")
            summary_var.set("")
            self._clear_dirty()
            refresh_history()

        tk.Button(form, text="提交并生成待审核凭证", font=FONT_B, bg=C_ACCENT, fg="white", relief="flat", padx=14, pady=7, command=submit).grid(row=2, column=6, columnspan=2, sticky="e", padx=4, pady=(10, 0))
        self._current_save = submit
        for variable in (date_var, amount_var, actor_var, summary_var, partner_var, bank_var, counterpart_var):
            variable.trace_add("write", lambda *_a: self._mark_dirty("收付款单"))
        refresh_history()

    # ------------------------------ 往来 ------------------------------
    def _page_partners(self, parent) -> None:
        self.header_status.configure(text="往来余额来自已记账凭证的辅助核算", fg=C_TEXT_MUTED)
        controls = tk.Frame(parent, bg=C_BG)
        controls.pack(fill="x", pady=(0, 10))
        type_var = tk.StringVar(value="全部")
        search_var = tk.StringVar()
        ttk.Combobox(controls, textvariable=type_var, values=("全部", "客户", "供应商"), width=10, state="readonly").pack(side="left")
        search = ttk.Entry(controls, textvariable=search_var, width=28)
        search.pack(side="left", padx=8)
        self._search_widget = search
        ttk.Button(controls, text="查找", command=lambda: refresh()).pack(side="left")
        tk.Button(controls, text="新增往来方", font=FONT_B, bg=C_NAV_ACTIVE, fg="white", relief="flat", padx=12, pady=6, command=lambda: self._add_partner_dialog(refresh)).pack(side="right")
        panes = ttk.Panedwindow(parent, orient="horizontal")
        panes.pack(fill="both", expand=True)
        left, right = self._panel(panes, padx=10, pady=10), self._panel(panes, padx=10, pady=10)
        panes.add(left, weight=2)
        panes.add(right, weight=3)
        partner_tree = self._tree(left, (("id", "ID", 45, "center"), ("type", "类型", 65, "center"), ("name", "名称", 180, "w"), ("ar", "应收", 100, "e"), ("ap", "应付", 100, "e")), height=16)
        tk.Label(right, text="往来明细", font=FONT_B, bg=C_PANEL, fg=C_TEXT).pack(anchor="w", pady=(0, 8))
        detail = self._tree(right, (("date", "日期", 90, "center"), ("vno", "凭证号", 140, "w"), ("summary", "摘要", 210, "w"), ("debit", "借方", 95, "e"), ("credit", "贷方", 95, "e")), height=16)
        partners: dict[str, dict] = {}

        def refresh():
            partner_tree.delete(*partner_tree.get_children())
            detail.delete(*detail.get_children())
            partners.clear()
            keyword = search_var.get().strip().lower()
            for row in arap.list_partners(None if type_var.get() == "全部" else type_var.get()):
                if keyword and keyword not in (str(row.get("name", "")) + str(row.get("tax_no", ""))).lower():
                    continue
                ar_value, ap_value = arap.partner_balance(row["id"])
                ar_cents = _parse_cents(abs(ar_value)) * (-1 if ar_value < 0 else 1)
                ap_cents = _parse_cents(abs(ap_value)) * (-1 if ap_value < 0 else 1)
                iid = str(row["id"])
                partners[iid] = row
                partner_tree.insert("", "end", iid=iid, values=(row["id"], row.get("type"), row.get("name"), _money(ar_cents), _money(ap_cents)))

        def selected(_event=None):
            detail.delete(*detail.get_children())
            selection = partner_tree.selection()
            if not selection:
                return
            partner = partners[selection[0]]
            for row in arap.ar_ap_ledger(partner["type"], partner["id"]):
                item = dict(row)
                detail.insert("", "end", values=(item.get("date"), item.get("vno"), item.get("vs") or item.get("summary"), _money(_row_cents(item, "debit_cents", "debit")), _money(_row_cents(item, "credit_cents", "credit"))))

        partner_tree.bind("<<TreeviewSelect>>", selected)
        search.bind("<Return>", lambda _e: refresh())
        type_var.trace_add("write", lambda *_a: self.root.after_idle(refresh))
        refresh()

    # ------------------------------ 费用票据 ------------------------------
    def _page_expense_invoice(self, parent) -> None:
        self.header_status.configure(text="审批、入账、付款分开处理；发票须人工确认", fg=C_TEXT_MUTED)
        notebook = ttk.Notebook(parent)
        notebook.pack(fill="both", expand=True)
        expense_tab, invoice_tab, retry_tab = tk.Frame(notebook, bg=C_PANEL), tk.Frame(notebook, bg=C_PANEL), tk.Frame(notebook, bg=C_PANEL)
        notebook.add(expense_tab, text="费用报销")
        notebook.add(invoice_tab, text="发票确认")
        notebook.add(retry_tab, text="失败重试")
        self._build_expense_tab(expense_tab)
        self._build_invoice_tab(invoice_tab)
        self._build_retry_tab(retry_tab)

    def _build_expense_tab(self, parent) -> None:
        controls = tk.Frame(parent, bg=C_PANEL, padx=12, pady=10)
        controls.pack(fill="x")
        ttk.Button(controls, text="新增员工", command=self._add_employee_dialog).pack(side="left")
        ttk.Button(controls, text="新增报销单", command=lambda: self._add_expense_dialog(refresh)).pack(side="left", padx=6)
        ttk.Button(controls, text="刷新", command=lambda: refresh()).pack(side="left")
        tree_panel = tk.Frame(parent, bg=C_PANEL, padx=12, pady=8)
        tree_panel.pack(fill="both", expand=True)
        tree = self._tree(tree_panel, (("id", "ID", 50, "center"), ("no", "单号", 130, "w"), ("employee", "员工", 90, "w"), ("date", "日期", 95, "center"), ("amount", "金额", 105, "e"), ("status", "状态", 90, "center"), ("approver", "审批人", 100, "w"), ("voucher", "凭证ID", 75, "center")), height=14)
        records: dict[str, dict] = {}

        def refresh():
            tree.delete(*tree.get_children())
            records.clear()
            for row in expense_mod.list_expense():
                item = dict(row)
                iid = str(item["id"])
                records[iid] = item
                tree.insert("", "end", iid=iid, values=(item["id"], item.get("no"), item.get("emp_name"), item.get("date"), _money(_row_cents(item, "amount_cents", "amount")), _status(item.get("status")), item.get("approver") or "", item.get("voucher_id") or ""))

        def current() -> dict | None:
            selection = tree.selection()
            return records.get(selection[0]) if selection else None

        def approve():
            row = current()
            if not row:
                messagebox.showwarning("请选择报销单", "请先选择一张报销单。", parent=self.root)
                return
            dialog = WorkflowDialog(self.root, "审核报销", self.default_actor)
            if not dialog.result:
                return
            try:
                result = _call_supported(expense_mod.approve, eid=row["id"], expense_id=row["id"], approver=dialog.result["actor"], actor=dialog.result["actor"], exception_reason=dialog.result["reason"] or None)
                ok, message, _data = _result(result)
                if not ok:
                    raise ValueError(message)
            except Exception as exc:
                messagebox.showerror("审核失败", str(exc), parent=self.root)
                return
            messagebox.showinfo("审核结果", message, parent=self.root)
            refresh()

        def send_feishu():
            row = current()
            if not row:
                messagebox.showwarning("请选择报销单", "请先选择一张报销单。", parent=self.root)
                return
            try:
                ok, message, _data = _result(feishu.submit_expense_approval(row))
                if not ok:
                    raise ValueError(message)
            except Exception as exc:
                messagebox.showerror("提交飞书失败", str(exc), parent=self.root)
                return
            messagebox.showinfo("飞书审批", message, parent=self.root)
            refresh()

        def sync_feishu():
            try:
                ok, message, _data = _result(feishu.sync_approval_status())
                if not ok:
                    raise ValueError(message)
            except Exception as exc:
                messagebox.showerror("同步失败", str(exc), parent=self.root)
                return
            messagebox.showinfo("飞书审批", message, parent=self.root)
            refresh()

        actions = tk.Frame(parent, bg=C_PANEL, padx=12, pady=9)
        actions.pack(fill="x")
        ttk.Button(actions, text="同步飞书状态", command=sync_feishu).pack(side="right", padx=4)
        ttk.Button(actions, text="提交飞书审批", command=send_feishu).pack(side="right", padx=4)
        ttk.Button(actions, text="本地审核", command=approve).pack(side="right", padx=4)
        refresh()

    def _build_invoice_tab(self, parent) -> None:
        controls = tk.Frame(parent, bg=C_PANEL, padx=12, pady=10)
        controls.pack(fill="x")
        status_var = tk.StringVar(value="全部")
        ttk.Combobox(controls, textvariable=status_var, values=("全部", "待确认", "草稿", "待审核", "已记账"), width=10, state="readonly").pack(side="left")
        ttk.Button(controls, text="刷新", command=lambda: refresh()).pack(side="left", padx=6)
        tk.Button(controls, text="运行发票识别", font=FONT_B, bg=C_NAV_ACTIVE, fg="white", relief="flat", padx=12, pady=6, command=self._fetch_invoices).pack(side="right")
        self.invoice_status = tk.Label(controls, text="", font=FONT_SMALL, bg=C_PANEL, fg=C_TEXT_MUTED)
        self.invoice_status.pack(side="right", padx=10)
        panel = tk.Frame(parent, bg=C_PANEL, padx=12, pady=8)
        panel.pack(fill="both", expand=True)
        tk.Label(panel, text="提示：发票须人工确认。生成凭证前请先到「往来」页建好销售方（供应商），再回来选发票点“确认并生成待审核凭证”。",
                 bg=C_PANEL, fg=C_TEXT_MUTED, font=FONT_SMALL, anchor="w").pack(fill="x", pady=(0, 6))
        tree = self._tree(panel, (("id", "ID", 45, "center"), ("number", "发票号", 150, "w"), ("date", "日期", 95, "center"), ("seller", "销售方", 210, "w"), ("total", "价税合计", 105, "e"), ("payment", "付款状态", 85, "center"), ("status", "状态", 85, "center"), ("voucher", "凭证ID", 70, "center")), height=14)
        records: dict[str, dict] = {}

        def refresh():
            tree.delete(*tree.get_children())
            records.clear()
            try:
                rows = invoice_gui.list_invoices(None)
            except Exception:
                rows = []
            for item in rows:
                row = dict(item)
                if status_var.get() != "全部" and _status(row.get("status")) != status_var.get():
                    continue
                iid = str(row.get("id"))
                records[iid] = row
                tree.insert("", "end", iid=iid, values=(row.get("id"), row.get("invoice_no"), row.get("date"), row.get("seller"), _money(_row_cents(row, "total_cents", "total")), row.get("payment_status") or "未知", _status(row.get("status")), row.get("voucher_id") or ""))

        def confirm():
            selection = tree.selection()
            if not selection:
                messagebox.showwarning("请选择发票", "请先选择一张发票。", parent=self.root)
                return
            if not arap.list_partners("供应商"):
                messagebox.showwarning("尚未建立供应商", "该发票生成凭证前需指定往来供应商。\n请先到「往来」页新增销售方（供应商），再回来确认。", parent=self.root)
                return
            row = records[selection[0]]
            dialog = InvoiceConfirmDialog(self.root, self._accounts(category="损益", direction="借"), arap.list_partners("供应商"), self.default_actor)
            if not dialog.result:
                return
            try:
                result = invoice_gui.invoice_to_voucher(
                    int(row["id"]),
                    expense_account=dialog.result["expense_account"],
                    partner_type="供应商",
                    partner_id=dialog.result["partner_id"],
                    payment_status=dialog.result["payment_status"],
                    deductible_tax=dialog.result["deductible_tax"],
                    creator=dialog.result["actor"],
                    submit=True,
                )
                ok, message, _data = _result(result)
                if not ok:
                    raise ValueError(message)
            except Exception as exc:
                messagebox.showerror("生成凭证失败", str(exc), parent=self.root)
                return
            messagebox.showinfo("发票已确认", message, parent=self.root)
            refresh()

        actions = tk.Frame(parent, bg=C_PANEL, padx=12, pady=9)
        actions.pack(fill="x")
        tk.Label(actions, text="未知付款状态默认记入应付账款；不会直接减少银行存款。", bg=C_PANEL, fg=C_TEXT_MUTED, font=FONT_SMALL).pack(side="left")
        tk.Button(actions, text="确认并生成待审核凭证", font=FONT_B, bg=C_ACCENT, fg="white", relief="flat", padx=12, pady=6, command=confirm).pack(side="right")
        status_var.trace_add("write", lambda *_a: self.root.after_idle(refresh))
        refresh()
        self._invoice_refresh = refresh

    def _build_retry_tab(self, parent: tk.Frame) -> None:
        import app_paths
        from modules.mail_processor import process_file

        controls = tk.Frame(parent, bg=C_PANEL, padx=12, pady=10)
        controls.pack(fill="x")
        tk.Label(controls, text="OCR 识别失败的文件，可在此重新处理。", bg=C_PANEL, fg=C_TEXT_MUTED, font=FONT_SMALL).pack(side="left")
        tk.Button(controls, text="刷新列表", command=lambda: self._retry_refresh() if hasattr(self, '_retry_refresh') else None).pack(side="right", padx=6)

        tree_panel = tk.Frame(parent, bg=C_PANEL, padx=12, pady=8)
        tree_panel.pack(fill="both", expand=True)
        cols = (("filename", "文件名", 260, "w"), ("size", "大小", 80, "e"), ("mtime", "修改时间", 150, "w"))
        tree = self._tree(tree_panel, cols, height=12)
        records: dict[str, Path] = {}

        def load_failed() -> None:
            tree.delete(*tree.get_children())
            records.clear()
            failed_dir = getattr(app_paths, "INVOICES_FAILED_DIR", None)
            if not failed_dir or not failed_dir.exists():
                return
            for path in sorted(failed_dir.iterdir(), key=lambda p: p.stat().st_mtime, reverse=True):
                if not path.is_file():
                    continue
                iid = str(path)
                records[iid] = path
                try:
                    size = path.stat().st_size
                    size_str = f"{size / 1024:.1f} KB" if size < 1024 * 1024 else f"{size / (1024 * 1024):.1f} MB"
                    mtime = dt.datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y-%m-%d %H:%M:%S")
                except OSError:
                    size_str = "?"
                    mtime = "?"
                tree.insert("", "end", iid=iid, values=(path.name, size_str, mtime))

        def retry_selected() -> None:
            selection = tree.selection()
            if not selection:
                messagebox.showwarning("请选择文件", "请先选择一个失败文件。", parent=self.root)
                return
            path = records[selection[0]]
            status_label.configure(text=f"正在处理: {path.name} ...", fg=C_WARN)
            parent.update()

            def worker():
                try:
                    result = process_file(str(path))
                    ok, message = result.ok, result.message
                    if ok:
                        try:
                            path.unlink()
                            LOG.info("重试成功，删除失败文件: %s", path)
                        except OSError:
                            pass
                except Exception as exc:
                    ok, message = False, str(exc)
                self.root.after(0, lambda: finished(ok, message))

            def finished(ok: bool, message: str):
                if not self.root.winfo_exists():
                    return
                status_label.configure(text=message, fg=C_OK if ok else C_DANGER)
                load_failed()

            threading.Thread(target=worker, daemon=True).start()

        actions = tk.Frame(parent, bg=C_PANEL, padx=12, pady=9)
        actions.pack(fill="x")
        tk.Label(actions, text="选中文件后点击「重试」；成功将自动删除原失败文件。", bg=C_PANEL, fg=C_TEXT_MUTED, font=FONT_SMALL).pack(side="left")
        tk.Button(actions, text="重试选中文件", font=FONT_B, bg=C_ACCENT, fg="white", relief="flat", padx=12, pady=6, command=retry_selected).pack(side="right")
        status_label = tk.Label(actions, text="", bg=C_PANEL, fg=C_TEXT_MUTED, font=FONT_SMALL)
        status_label.pack(side="right", padx=10)

        load_failed()
        self._retry_refresh = load_failed

    def _fetch_invoices(self) -> None:
        if not hasattr(self, "invoice_status"):
            return
        self.invoice_status.configure(text="正在识别，请稍候…", fg=C_WARN)

        def worker():
            try:
                result = invoice_gui.fetch_and_recognize()
                ok, message, data = _result(result)
                if ok and data:
                    self._invoice_summary_path = str(data)
                    for raw in invoice_gui.load_invoices(data):
                        invoice_gui.register_invoice(raw, raw.get("file_path") or raw.get("file"))
            except Exception as exc:
                ok, message = False, str(exc)
            self.root.after(0, lambda: finished(ok, message))

        def finished(ok: bool, message: str):
            if not self.root.winfo_exists():
                return
            self.invoice_status.configure(text=message, fg=C_OK if ok else C_DANGER)
            if ok and callable(getattr(self, "_invoice_refresh", None)):
                self._invoice_refresh()

        threading.Thread(target=worker, daemon=True).start()

    # ------------------------------ 邮件收票 ------------------------------
    def _page_mail(self, parent) -> None:
        from gui.page_mail import MailPage

        page = MailPage(parent, self)
        page.pack(fill="both", expand=True)

    # ------------------------------ 进销存 ------------------------------
    def _page_inventory(self, parent) -> None:
        from gui.page_inventory import InventoryPage

        page = InventoryPage(parent, self)
        page.pack(fill="both", expand=True)

    # ------------------------------ 期末处理 ------------------------------
    def _page_closing(self, parent) -> None:
        from gui.page_closing import ClosingPage

        page = ClosingPage(parent, self)
        page.pack(fill="both", expand=True)

    # ------------------------------ 报表 ------------------------------
    def _page_reports(self, parent) -> None:
        self.header_status.configure(text="正式报表仅使用已记账凭证，并可追溯至明细", fg=C_TEXT_MUTED)
        controls = tk.Frame(parent, bg=C_BG)
        controls.pack(fill="x", pady=(0, 10))
        period_var = tk.StringVar(value=_current_period())
        tk.Label(controls, text="会计期间", bg=C_BG, font=FONT).pack(side="left")
        ttk.Combobox(controls, textvariable=period_var, values=self._period_values(), width=12, state="readonly").pack(side="left", padx=6)
        generate_btn = ttk.Button(controls, text="生成报表")
        export_btn = ttk.Button(controls, text="导出 CSV", command=lambda: _do_export("csv"))
        excel_btn = ttk.Button(controls, text="导出 Excel")
        pdf_btn = ttk.Button(controls, text="导出 PDF")
        generate_btn.pack(side="left", padx=4)
        export_btn.pack(side="left", padx=4)
        excel_btn.pack(side="left", padx=4)
        pdf_btn.pack(side="left", padx=4)
        self._report_action_buttons = [generate_btn, export_btn, excel_btn, pdf_btn]
        if self._issues:
            for button in self._report_action_buttons:
                button.configure(state="disabled")
            tk.Label(controls, text="请先处理阻断性数据异常", bg=C_BG, fg=C_DANGER, font=FONT_B).pack(side="right")

        notebook = ttk.Notebook(parent)
        notebook.pack(fill="both", expand=True)
        bs_tab, is_tab, cf_tab, ab_tab, journal_tab = [tk.Frame(notebook, bg=C_PANEL) for _ in range(5)]
        notebook.add(bs_tab, text="资产负债表")
        notebook.add(is_tab, text="利润表")
        notebook.add(cf_tab, text="现金流量表")
        notebook.add(ab_tab, text="科目余额表")
        notebook.add(journal_tab, text="序时簿")
        bs_tree = self._tree(bs_tab, (("kind", "类别", 80, "center"), ("code", "项目/科目", 120, "w"), ("name", "名称", 270, "w"), ("amount", "期末金额", 140, "e")), height=17)
        is_tree = self._tree(is_tab, (("kind", "类别", 80, "center"), ("code", "科目", 100, "w"), ("name", "名称", 270, "w"), ("month", "本月数", 130, "e"), ("ytd", "本年累计", 130, "e")), height=17)
        cf_tree = self._tree(cf_tab, (("name", "项目", 340, "w"), ("amount", "金额", 150, "e")), height=12)
        ab_tree = self._tree(ab_tab, (("code", "科目编码", 95, "w"), ("name", "科目名称", 170, "w"), ("opening", "期初余额", 110, "e"), ("debit", "本期借方", 110, "e"), ("credit", "本期贷方", 110, "e"), ("ending_debit", "期末借方", 110, "e"), ("ending_credit", "期末贷方", 110, "e")), height=15)
        journal_tree = self._tree(journal_tab, (("date", "日期", 95, "center"), ("vno", "凭证号", 135, "w"), ("seq", "序号", 55, "center"), ("account", "科目", 170, "w"), ("summary", "摘要", 260, "w"), ("debit", "借方", 110, "e"), ("credit", "贷方", 110, "e"), ("source", "来源", 85, "center")), height=15)
        footer = tk.Label(parent, text="", font=FONT_B, bg=C_BG, fg=C_TEXT)
        footer.pack(anchor="e", pady=(8, 0))

        def generate():
            if self._open_issues():
                self._refresh_issue_state()
                messagebox.showerror("报表已暂停", "存在未处理的数据异常，请先处理后再生成正式报表。", parent=self.root)
                return
            try:
                period = _validate_period(period_var.get())
                if hasattr(reports, "generate_report_bundle"):
                    bundle = reports.generate_report_bundle(period)
                else:
                    bundle = {"period": period, "balance_sheet": reports.balance_sheet(period), "income_statement": reports.income_statement(period), "cash_flow_statement": reports.cash_flow_statement(period)}
            except Exception as exc:
                if isinstance(exc, getattr(reports, "ReportBlockedError", RuntimeError)):
                    self._refresh_issue_state()
                messagebox.showerror("生成失败", str(exc), parent=self.root)
                return
            self._report_cache = bundle
            bs_tree.delete(*bs_tree.get_children())
            is_tree.delete(*is_tree.get_children())
            cf_tree.delete(*cf_tree.get_children())
            ab_tree.delete(*ab_tree.get_children())
            journal_tree.delete(*journal_tree.get_children())
            bs = bundle["balance_sheet"]
            for kind, key in (("资产", "assets"), ("负债", "liabilities"), ("权益", "equity")):
                for row in bs.get(key, []):
                    bs_tree.insert("", "end", values=(kind, row.get("code"), row.get("name"), _money(_row_cents(row, "amount_cents", "amount"))))
            bs_tree.insert("", "end", values=("合计", "", "资产合计", _money(_row_cents(bs, "total_assets_cents", "total_assets"))))
            bs_tree.insert("", "end", values=("合计", "", "负债和权益合计", _money(_row_cents(bs, "total_liabilities_cents", "total_liabilities") + _row_cents(bs, "total_equity_cents", "total_equity"))))
            income = bundle["income_statement"]
            current, ytd = income.get("current", income), income.get("year_to_date", income)
            current_rows = {row.get("code"): row for key in ("revenue", "expense") for row in current.get(key, [])}
            ytd_rows = {row.get("code"): row for key in ("revenue", "expense") for row in ytd.get(key, [])}
            codes = list(dict.fromkeys([*current_rows, *ytd_rows]))
            for code in codes:
                row = current_rows.get(code) or ytd_rows.get(code)
                kind = "收入" if row in current.get("revenue", []) or row in ytd.get("revenue", []) else "费用"
                month_value = (current_rows.get(code) or {}).get("amount_cents", 0)
                ytd_value = (ytd_rows.get(code) or {}).get("amount_cents", 0)
                is_tree.insert("", "end", values=(kind, code, row.get("name"), _money(month_value), _money(ytd_value)))
            is_tree.insert("", "end", values=("合计", "", "净利润", _money(current.get("net_profit_cents", 0)), _money(ytd.get("net_profit_cents", 0))))
            cash = bundle["cash_flow_statement"]
            cash_rows = (("期初现金余额", "opening_cash_cents"), ("经营活动流入", "operating_in_cents"), ("经营活动流出", "operating_out_cents"), ("经营活动净额", "operating_net_cents"), ("投资活动净额", "investing_cents"), ("筹资活动净额", "financing_cents"), ("未分类差额", "unclassified_cents"), ("现金净增加额", "net_cash_cents"), ("期末现金余额", "ending_cash_cents"))
            for label, key in cash_rows:
                cf_tree.insert("", "end", values=(label, _money(cash.get(key, 0))))
            account_table = bundle.get("account_balance_table") or reports.account_balance_table(period)
            for row in account_table.get("rows", []):
                ab_tree.insert("", "end", values=(
                    row.get("code"), row.get("name"),
                    _money(_row_cents(row, "opening_debit_cents", "opening_debit") - _row_cents(row, "opening_credit_cents", "opening_credit")),
                    _money(_row_cents(row, "period_debit_cents", "period_debit")),
                    _money(_row_cents(row, "period_credit_cents", "period_credit")),
                    _money(_row_cents(row, "ending_debit_cents", "ending_debit")),
                    _money(_row_cents(row, "ending_credit_cents", "ending_credit")),
                ))
            journal = bundle.get("chronological_ledger") or reports.chronological_ledger(period)
            for row in journal.get("rows", []):
                account = f"{row.get('account_code') or ''} {row.get('account_name') or ''}".strip()
                journal_tree.insert("", "end", values=(
                    row.get("date"), row.get("vno"), row.get("seq"), account,
                    row.get("line_summary") or row.get("voucher_summary") or "",
                    _money(_row_cents(row, "debit_cents", "debit")),
                    _money(_row_cents(row, "credit_cents", "credit")),
                    row.get("source") or "",
                ))
            bs_ok = bool(bs.get("balanced"))
            cf_ok = bool(cash.get("reconciled"))
            journal_ok = bool(journal.get("balanced"))
            footer.configure(text=f"资产负债表 {'平衡' if bs_ok else '不平衡'} · 现金流量 {'勾稽一致' if cf_ok else '存在差额'} · 序时簿 {'借贷相等' if journal_ok else '借贷不等'}", fg=C_OK if bs_ok and cf_ok and journal_ok else C_DANGER)

        def _do_export(fmt: str):
            if not self._report_cache or self._report_cache.get("period") != period_var.get():
                generate()
            if not self._report_cache:
                return
            ext = {"csv": ".csv", "xlsx": ".xlsx", "pdf": ".pdf"}.get(fmt, ".csv")
            initial_dir = str(getattr(app_paths, "REPORTS_DIR", ROOT / "data" / "reports"))
            path = filedialog.asksaveasfilename(parent=self.root, title="导出财务报表", initialdir=initial_dir, initialfile=f"财务报表_{period_var.get()}{ext}", defaultextension=ext, filetypes=((f"{fmt.upper()} 文件", f"*{ext}"), ("所有文件", "*.*")))
            if not path:
                return
            try:
                if fmt == "csv":
                    self._export_report_csv(Path(path), self._report_cache)
                elif fmt == "xlsx":
                    _export_report_xlsx(Path(path), self._report_cache)
                elif fmt == "pdf":
                    _export_report_pdf(Path(path), self._report_cache)
                else:
                    messagebox.showerror("导出失败", f"不支持的格式: {fmt}", parent=self.root)
                    return
            except Exception as exc:
                messagebox.showerror("导出失败", str(exc), parent=self.root)
                return
            messagebox.showinfo("导出完成", f"报表已导出到：\n{path}", parent=self.root)

        def drill(tree: ttk.Treeview):
            selection = tree.selection()
            if not selection:
                return
            values = tree.item(selection[0], "values")
            code = str(values[1]) if len(values) > 1 else ""
            if re.fullmatch(r"\d{4}", code):
                self._pending_ledger_account = (code, period_var.get())
                self._show_page("账务")

        bs_tree.bind("<Double-1>", lambda _e: drill(bs_tree))
        is_tree.bind("<Double-1>", lambda _e: drill(is_tree))
        generate_btn.configure(command=generate)
        export_btn.configure(command=lambda: _do_export("csv"))
        excel_btn.configure(command=lambda: _do_export("xlsx"))
        pdf_btn.configure(command=lambda: _do_export("pdf"))
        if not self._issues:
            self.root.after_idle(generate)

    # ------------------------------ 账簿 ------------------------------
    def _page_books(self, parent) -> None:
        self.header_status.configure(text="账簿：科目余额表 + 序时簿，仅取已记账凭证", fg=C_TEXT_MUTED)
        controls = tk.Frame(parent, bg=C_BG)
        controls.pack(fill="x", pady=(0, 10))
        period_var = tk.StringVar(value=_latest_posted_period() or _current_period())
        tk.Label(controls, text="会计期间", bg=C_BG, font=FONT).pack(side="left")
        ttk.Combobox(controls, textvariable=period_var, values=self._period_values(), width=12, state="readonly").pack(side="left", padx=6)
        zero_var = tk.BooleanVar(value=False)
        tk.Checkbutton(controls, text="显示零余额科目", variable=zero_var, bg=C_BG, font=FONT, activebackground=C_BG, selectcolor=C_CARD).pack(side="left", padx=8)
        gen_btn = ttk.Button(controls, text="生成账簿")
        csv_btn = ttk.Button(controls, text="导出 CSV")
        xlsx_btn = ttk.Button(controls, text="导出 Excel")
        pdf_btn = ttk.Button(controls, text="导出 PDF")
        for button in (gen_btn, csv_btn, xlsx_btn, pdf_btn):
            button.pack(side="left", padx=4)

        notebook = ttk.Notebook(parent)
        notebook.pack(fill="both", expand=True)
        ab_tab, journal_tab = [tk.Frame(notebook, bg=C_PANEL) for _ in range(2)]
        notebook.add(ab_tab, text="科目余额表")
        notebook.add(journal_tab, text="序时簿")
        ab_tree = self._tree(ab_tab, (("code", "科目编码", 95, "w"), ("name", "科目名称", 180, "w"), ("category", "类别", 70, "center"), ("opening", "期初余额", 120, "e"), ("debit", "本期借方", 120, "e"), ("credit", "本期贷方", 120, "e"), ("ending_debit", "期末借方", 120, "e"), ("ending_credit", "期末贷方", 120, "e")), height=15)
        journal_head = tk.Frame(journal_tab, bg=C_PANEL)
        journal_head.pack(fill="x", pady=(6, 2))
        tk.Label(journal_head, text="科目", bg=C_PANEL, font=FONT).pack(side="left", padx=(6, 4))
        account_var = tk.StringVar(value=_ALL_ACCOUNTS)
        account_box = ttk.Combobox(journal_head, textvariable=account_var, values=[_ALL_ACCOUNTS], width=24, state="readonly")
        account_box.pack(side="left")
        ledger_btn = ttk.Button(journal_head, text="在账务页看明细账 →")
        ledger_btn.pack(side="left", padx=8)
        tk.Label(journal_head, text="双击上方「科目余额表」任一科目，即在下方查看该科目的序时簿", bg=C_PANEL, font=FONT_SMALL, fg=C_TEXT_MUTED).pack(side="left", padx=6)
        journal_tree = self._tree(journal_tab, (("date", "日期", 95, "center"), ("vno", "凭证号", 135, "w"), ("seq", "序号", 55, "center"), ("account", "科目", 180, "w"), ("summary", "摘要", 250, "w"), ("debit", "借方", 110, "e"), ("credit", "贷方", 110, "e"), ("balance", "余额", 120, "e"), ("source", "来源", 80, "center")), height=15)
        footer = tk.Label(parent, text="", font=FONT_B, bg=C_BG, fg=C_TEXT)
        footer.pack(anchor="e", pady=(8, 0))
        cache: dict = {"period": None}
        journal_rows: list = []

        def _ab_values(row: dict) -> tuple:
            opening = _row_cents(row, "opening_debit_cents", "opening_debit") - _row_cents(row, "opening_credit_cents", "opening_credit")
            return (
                row.get("code"), row.get("name"), row.get("category") or "",
                _money(opening),
                _money(_row_cents(row, "period_debit_cents", "period_debit")),
                _money(_row_cents(row, "period_credit_cents", "period_credit")),
                _money(_row_cents(row, "ending_debit_cents", "ending_debit")),
                _money(_row_cents(row, "ending_credit_cents", "ending_credit")),
            )

        def _render_journal():
            # 依「科目」过滤框渲染序时簿；选中单科目时余额列=该科目逐笔累计（三栏式明细账）
            journal_tree.delete(*journal_tree.get_children())
            chosen = account_var.get()
            code = "" if chosen == _ALL_ACCOUNTS else chosen.split(" ", 1)[0]
            opening = 0
            for row in (cache.get("account_balance_table") or {}).get("rows", []):
                if code and str(row.get("code")) == code:
                    opening = _row_cents(row, "opening_debit_cents", "opening_debit") - _row_cents(row, "opening_credit_cents", "opening_credit")
                    break
            running = opening
            shown: list = []
            for row in journal_rows:
                if code and str(row.get("account_code")) != code:
                    continue
                if code:
                    running += _row_cents(row, "debit_cents", "debit") - _row_cents(row, "credit_cents", "credit")
                    balance_cents = running
                else:
                    balance_cents = _row_cents(row, "running_debit_cents", "running_debit") - _row_cents(row, "running_credit_cents", "running_credit")
                account = f"{row.get('account_code') or ''} {row.get('account_name') or ''}".strip()
                journal_tree.insert("", "end", values=(
                    row.get("date"), row.get("vno"), row.get("seq"), account,
                    row.get("line_summary") or row.get("voucher_summary") or "",
                    _money(_row_cents(row, "debit_cents", "debit")),
                    _money(_row_cents(row, "credit_cents", "credit")),
                    _money(balance_cents),
                    row.get("source") or "",
                ))
                shown.append(row)
            full = cache.get("chronological_ledger") or {}
            cache["chronological_ledger"] = {**full, "rows": shown}
            balanced = bool(full.get("balanced"))
            scope = _ALL_ACCOUNTS if not code else chosen
            footer.configure(
                text=f"科目余额表 {cache.get('ab_count', 0)} 行 · 序时簿[{scope}] {len(shown)} 行 · 借贷{'相等' if balanced else '不等'}",
                fg=C_OK if balanced else C_DANGER,
            )

        def generate():
            try:
                period = _validate_period(period_var.get())
                account_table = reports.account_balance_table(period, include_zero=zero_var.get())
                journal = reports.chronological_ledger(period=period)
            except Exception as exc:
                if isinstance(exc, getattr(reports, "ReportBlockedError", RuntimeError)):
                    self._refresh_issue_state()
                messagebox.showerror("生成失败", str(exc), parent=self.root)
                return
            cache.clear()
            journal_rows.clear()
            journal_rows.extend(journal.get("rows", []))
            cache.update({"period": period, "account_balance_table": account_table, "chronological_ledger": journal})
            options = [f"{row.get('code')} {row.get('name')}".strip() for row in account_table.get("rows", [])]
            options.sort()
            account_box.configure(values=[_ALL_ACCOUNTS] + options)
            if account_var.get() not in ([_ALL_ACCOUNTS] + options):
                account_var.set(_ALL_ACCOUNTS)
            ab_tree.delete(*ab_tree.get_children())
            for row in account_table.get("rows", []):
                ab_tree.insert("", "end", values=_ab_values(row))
            cache["ab_count"] = len(account_table.get("rows", []))
            _render_journal()

        def _do_export(fmt: str):
            if cache.get("period") != period_var.get():
                generate()
                if cache.get("period") != period_var.get():
                    return
            ext = {"csv": ".csv", "xlsx": ".xlsx", "pdf": ".pdf"}.get(fmt, ".csv")
            initial_dir = str(getattr(app_paths, "REPORTS_DIR", ROOT / "data" / "reports"))
            path = filedialog.asksaveasfilename(parent=self.root, title="导出账簿", initialdir=initial_dir, initialfile=f"账簿_{period_var.get()}{ext}", defaultextension=ext, filetypes=((f"{fmt.upper()} 文件", f"*{ext}"), ("所有文件", "*.*")))
            if not path:
                return
            try:
                if fmt == "csv":
                    self._export_books_csv(Path(path), cache)
                elif fmt == "xlsx":
                    self._export_books_xlsx(Path(path), cache)
                elif fmt == "pdf":
                    self._export_books_pdf(Path(path), cache)
                else:
                    messagebox.showerror("导出失败", f"不支持的格式: {fmt}", parent=self.root)
                    return
            except Exception as exc:
                messagebox.showerror("导出失败", str(exc), parent=self.root)
                return
            messagebox.showinfo("导出完成", f"账簿已导出到：\n{path}", parent=self.root)

        def drill(_event=None):
            selection = ab_tree.selection()
            if not selection:
                return
            code = str(ab_tree.item(selection[0], "values")[0] or "").strip()
            if not code:
                return
            target = next((str(v) for v in account_box.cget("values") if str(v).split(" ", 1)[0] == code), None)
            if target:
                account_var.set(target)
                notebook.select(journal_tab)
                _render_journal()

        def open_ledger():
            chosen = account_var.get()
            if chosen == _ALL_ACCOUNTS:
                messagebox.showinfo("提示", "请先在序时簿里选择一个科目。", parent=self.root)
                return
            self._pending_ledger_account = (chosen.split(" ", 1)[0], period_var.get())
            self._show_page("账务")

        ab_tree.bind("<Double-1>", drill)
        account_box.bind("<<ComboboxSelected>>", lambda _e: _render_journal())
        ledger_btn.configure(command=open_ledger)
        gen_btn.configure(command=generate)
        csv_btn.configure(command=lambda: _do_export("csv"))
        xlsx_btn.configure(command=lambda: _do_export("xlsx"))
        pdf_btn.configure(command=lambda: _do_export("pdf"))
        if not self._issues:
            self.root.after_idle(generate)

    def _export_books_csv(self, path: Path, cache: dict) -> None:
        period = cache.get("period") or ""
        account_table = cache.get("account_balance_table") or {}
        journal = cache.get("chronological_ledger") or {}
        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(["科目余额表", period])
            writer.writerow(["科目编码", "科目名称", "类别", "期初余额", "本期借方", "本期贷方", "期末借方", "期末贷方"])
            for row in account_table.get("rows", []):
                opening = _row_cents(row, "opening_debit_cents", "opening_debit") - _row_cents(row, "opening_credit_cents", "opening_credit")
                writer.writerow((
                    row.get("code"), row.get("name"), row.get("category") or "", _money(opening),
                    _money(_row_cents(row, "period_debit_cents", "period_debit")),
                    _money(_row_cents(row, "period_credit_cents", "period_credit")),
                    _money(_row_cents(row, "ending_debit_cents", "ending_debit")),
                    _money(_row_cents(row, "ending_credit_cents", "ending_credit")),
                ))
            writer.writerow([])
            writer.writerow(["序时簿", journal.get("start_period", period), journal.get("end_period", period)])
            writer.writerow(["日期", "凭证号", "序号", "科目编码", "科目名称", "摘要", "借方", "贷方", "来源"])
            for row in journal.get("rows", []):
                writer.writerow((
                    row.get("date"), row.get("vno"), row.get("seq"), row.get("account_code"),
                    row.get("account_name"), row.get("line_summary") or row.get("voucher_summary") or "",
                    _money(_row_cents(row, "debit_cents", "debit")),
                    _money(_row_cents(row, "credit_cents", "credit")), row.get("source") or "",
                ))

    def _export_books_xlsx(self, path: Path, cache: dict) -> None:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "科目余额表"
        ws.append(["科目编码", "科目名称", "类别", "期初余额", "本期借方", "本期贷方", "期末借方", "期末贷方"])
        for row in (cache.get("account_balance_table") or {}).get("rows", []):
            opening = _row_cents(row, "opening_debit_cents", "opening_debit") - _row_cents(row, "opening_credit_cents", "opening_credit")
            ws.append((
                row.get("code"), row.get("name"), row.get("category") or "", _money(opening),
                _money(_row_cents(row, "period_debit_cents", "period_debit")),
                _money(_row_cents(row, "period_credit_cents", "period_credit")),
                _money(_row_cents(row, "ending_debit_cents", "ending_debit")),
                _money(_row_cents(row, "ending_credit_cents", "ending_credit")),
            ))
        ws2 = wb.create_sheet("序时簿")
        ws2.append(["日期", "凭证号", "序号", "科目编码", "科目名称", "摘要", "借方", "贷方", "来源"])
        for row in (cache.get("chronological_ledger") or {}).get("rows", []):
            ws2.append((
                row.get("date"), row.get("vno"), row.get("seq"), row.get("account_code"), row.get("account_name"),
                row.get("line_summary") or row.get("voucher_summary") or "",
                _money(_row_cents(row, "debit_cents", "debit")),
                _money(_row_cents(row, "credit_cents", "credit")), row.get("source") or "",
            ))
        for sheet in wb.worksheets:
            for column in sheet.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        max_length = max(max_length, len(str(cell.value or "")))
                    except Exception:
                        pass
                sheet.column_dimensions[column_letter].width = min(max_length + 2, 60)
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(path))

    def _export_books_pdf(self, path: Path, cache: dict) -> None:
        from fpdf import FPDF
        pdf = FPDF()
        cjk = self._resolve_cjk_font()
        font = "Helvetica"
        if cjk:
            try:
                pdf.add_font("CJK", "", cjk)
                pdf.add_font("CJK", "B", cjk)
                font = "CJK"
            except Exception:
                font = "Helvetica"
        pdf.add_page()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.set_font(font, "B", 14)
        pdf.cell(0, 10, f"账簿 {cache.get('period') or ''}", ln=True, align="L")
        pdf.ln(2)

        def _section(title, headers, rows):
            pdf.set_font(font, "B", 11)
            pdf.cell(0, 8, title, ln=True)
            pdf.set_font(font, "B", 8)
            col_w = 190 / max(len(headers), 1)
            for h in headers:
                pdf.cell(col_w, 7, str(h), border=1, align="C")
            pdf.ln()
            pdf.set_font(font, "", 8)
            for row in rows:
                for value in row:
                    pdf.cell(col_w, 7, str(value) if value is not None else "", border=1, align="L")
                pdf.ln()
            pdf.ln(3)

        ab_rows = []
        for row in (cache.get("account_balance_table") or {}).get("rows", []):
            opening = _row_cents(row, "opening_debit_cents", "opening_debit") - _row_cents(row, "opening_credit_cents", "opening_credit")
            ab_rows.append((
                row.get("code"), row.get("name"), _money(opening),
                _money(_row_cents(row, "period_debit_cents", "period_debit")),
                _money(_row_cents(row, "period_credit_cents", "period_credit")),
                _money(_row_cents(row, "ending_debit_cents", "ending_debit")),
                _money(_row_cents(row, "ending_credit_cents", "ending_credit")),
            ))
        _section("科目余额表", ["科目编码", "科目名称", "期初余额", "本期借方", "本期贷方", "期末借方", "期末贷方"], ab_rows)
        journal_rows = []
        for row in (cache.get("chronological_ledger") or {}).get("rows", []):
            journal_rows.append((
                row.get("date"), row.get("vno"), row.get("seq"), row.get("account_code"), row.get("account_name"),
                row.get("line_summary") or row.get("voucher_summary") or "",
                _money(_row_cents(row, "debit_cents", "debit")),
                _money(_row_cents(row, "credit_cents", "credit")), row.get("source") or "",
            ))
        _section("序时簿", ["日期", "凭证号", "序号", "科目编码", "科目名称", "摘要", "借方", "贷方", "来源"], journal_rows)
        path.parent.mkdir(parents=True, exist_ok=True)
        pdf.output(str(path))

    def _export_report_csv(self, path: Path, bundle: dict) -> None:
        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            period = bundle["period"]
            bs = bundle["balance_sheet"]
            writer.writerow(["资产负债表", period])
            writer.writerow(["类别", "项目/科目", "名称", "期末金额"])
            for kind, key in (("资产", "assets"), ("负债", "liabilities"), ("权益", "equity")):
                for row in bs.get(key, []):
                    writer.writerow((kind, row.get("code"), row.get("name"), _money(row.get("amount_cents", 0))))
            writer.writerow([])
            income = bundle["income_statement"]
            current, ytd = income.get("current", income), income.get("year_to_date", income)
            writer.writerow(["利润表", period])
            writer.writerow(["类别", "科目", "名称", "本月数", "本年累计"])
            cur = {row["code"]: row for key in ("revenue", "expense") for row in current.get(key, [])}
            year = {row["code"]: row for key in ("revenue", "expense") for row in ytd.get(key, [])}
            for code in dict.fromkeys([*cur, *year]):
                row = cur.get(code) or year[code]
                kind = "收入" if row in current.get("revenue", []) or row in ytd.get("revenue", []) else "费用"
                writer.writerow((kind, code, row.get("name"), _money((cur.get(code) or {}).get("amount_cents", 0)), _money((year.get(code) or {}).get("amount_cents", 0))))
            writer.writerow(("合计", "", "净利润", _money(current.get("net_profit_cents", 0)), _money(ytd.get("net_profit_cents", 0))))
            writer.writerow([])
            cash = bundle["cash_flow_statement"]
            writer.writerow(["现金流量表", period])
            for label, key in (("期初现金余额", "opening_cash_cents"), ("经营活动净额", "operating_net_cents"), ("投资活动净额", "investing_cents"), ("筹资活动净额", "financing_cents"), ("未分类差额", "unclassified_cents"), ("现金净增加额", "net_cash_cents"), ("期末现金余额", "ending_cash_cents")):
                writer.writerow((label, _money(cash.get(key, 0))))
            writer.writerow([])
            account_table = bundle.get("account_balance_table") or {}
            writer.writerow(["科目余额表", period])
            writer.writerow(["科目编码", "科目名称", "期初余额", "本期借方", "本期贷方", "期末借方", "期末贷方"])
            for row in account_table.get("rows", []):
                opening = _row_cents(row, "opening_debit_cents", "opening_debit") - _row_cents(row, "opening_credit_cents", "opening_credit")
                writer.writerow((
                    row.get("code"), row.get("name"), _money(opening),
                    _money(_row_cents(row, "period_debit_cents", "period_debit")),
                    _money(_row_cents(row, "period_credit_cents", "period_credit")),
                    _money(_row_cents(row, "ending_debit_cents", "ending_debit")),
                    _money(_row_cents(row, "ending_credit_cents", "ending_credit")),
                ))
            writer.writerow([])
            journal = bundle.get("chronological_ledger") or {}
            writer.writerow(["序时簿", journal.get("start_period", period), journal.get("end_period", period)])
            writer.writerow(["日期", "凭证号", "序号", "科目编码", "科目名称", "摘要", "借方", "贷方", "来源"])
            for row in journal.get("rows", []):
                writer.writerow((
                    row.get("date"), row.get("vno"), row.get("seq"), row.get("account_code"),
                    row.get("account_name"), row.get("line_summary") or row.get("voucher_summary") or "",
                    _money(_row_cents(row, "debit_cents", "debit")),
                    _money(_row_cents(row, "credit_cents", "credit")), row.get("source") or "",
                ))

    def _export_report_xlsx(self, path: Path, bundle: dict) -> None:
        from openpyxl import Workbook
        wb = Workbook()
        period = bundle["period"]
        bs = bundle["balance_sheet"]
        ws = wb.active
        ws.title = "资产负债表"
        ws.append(["类别", "项目/科目", "名称", "期末金额"])
        for kind, key in (("资产", "assets"), ("负债", "liabilities"), ("权益", "equity")):
            for row in bs.get(key, []):
                ws.append((kind, row.get("code"), row.get("name"), _money(row.get("amount_cents", 0))))
        ws.append([])
        income = bundle["income_statement"]
        current, ytd = income.get("current", income), income.get("year_to_date", income)
        ws2 = wb.create_sheet("利润表")
        ws2.append(["类别", "科目", "名称", "本月数", "本年累计"])
        cur = {row["code"]: row for key in ("revenue", "expense") for row in current.get(key, [])}
        year = {row["code"]: row for key in ("revenue", "expense") for row in ytd.get(key, [])}
        for code in dict.fromkeys([*cur, *year]):
            row = cur.get(code) or year[code]
            kind = "收入" if row in current.get("revenue", []) or row in ytd.get("revenue", []) else "费用"
            ws2.append((kind, code, row.get("name"), _money((cur.get(code) or {}).get("amount_cents", 0)), _money((year.get(code) or {}).get("amount_cents", 0))))
        ws2.append(("合计", "", "净利润", _money(current.get("net_profit_cents", 0)), _money(ytd.get("net_profit_cents", 0))))
        ws3 = wb.create_sheet("科目余额表")
        ws3.append(["科目编码", "科目名称", "期初余额", "本期借方", "本期贷方", "期末借方", "期末贷方"])
        account_table = bundle.get("account_balance_table") or {}
        for row in account_table.get("rows", []):
            opening = _row_cents(row, "opening_debit_cents", "opening_debit") - _row_cents(row, "opening_credit_cents", "opening_credit")
            ws3.append((row.get("code"), row.get("name"), _money(opening), _money(_row_cents(row, "period_debit_cents", "period_debit")), _money(_row_cents(row, "period_credit_cents", "period_credit")), _money(_row_cents(row, "ending_debit_cents", "ending_debit")), _money(_row_cents(row, "ending_credit_cents", "ending_credit"))))
        ws4 = wb.create_sheet("序时簿")
        ws4.append(["日期", "凭证号", "序号", "科目编码", "科目名称", "摘要", "借方", "贷方", "来源"])
        journal = bundle.get("chronological_ledger") or {}
        for row in journal.get("rows", []):
            ws4.append((row.get("date"), row.get("vno"), row.get("seq"), row.get("account_code"), row.get("account_name"), row.get("line_summary") or row.get("voucher_summary") or "", _money(_row_cents(row, "debit_cents", "debit")), _money(_row_cents(row, "credit_cents", "credit")), row.get("source") or ""))
        for sheet in wb.worksheets:
            for column in sheet.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        max_length = max(max_length, len(str(cell.value or "")))
                    except Exception:
                        pass
                sheet.column_dimensions[column_letter].width = min(max_length + 2, 60)
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(path))

    def _resolve_cjk_font(self) -> str | None:
        """优先使用随包内置的 CJK 字体；找不到时回退到系统 Windows 字体目录。"""
        here = Path(__file__).resolve().parent
        candidates = [
            here / "simhei.ttf",
            Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "simhei.ttf",
            Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "msyh.ttc",
        ]
        for c in candidates:
            if c.exists():
                return str(c)
        return None

    def _export_report_pdf(self, path: Path, bundle: dict) -> None:
        from fpdf import FPDF
        pdf = FPDF()
        cjk = self._resolve_cjk_font()
        if cjk:
            try:
                pdf.add_font("CJK", "", cjk)
                pdf.add_font("CJK", "B", cjk)
                pdf.add_font("CJK", "I", cjk)
                pdf.add_font("CJK", "BI", cjk)
                font = "CJK"
            except Exception:
                font = "Helvetica"
        else:
            font = "Helvetica"
        pdf.add_page()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.set_font(font, "B", 14)
        pdf.cell(0, 10, f"财务报表 {bundle['period']}", ln=True, align="L")
        pdf.ln(2)
        # Helper to write a section
        def _section(title, headers, rows):
            pdf.set_font(font, "B", 11)
            pdf.cell(0, 8, title, ln=True)
            pdf.set_font(font, "B", 9)
            col_w = 190 / max(len(headers), 1)
            for h in headers:
                pdf.cell(col_w, 7, str(h), border=1, align="C")
            pdf.ln()
            pdf.set_font(font, "", 9)
            for row in rows:
                for value in row:
                    pdf.cell(col_w, 7, str(value) if value is not None else "", border=1, align="L")
                pdf.ln()
            pdf.ln(3)
        bs = bundle["balance_sheet"]
        bs_rows = []
        for kind, key in (("资产", "assets"), ("负债", "liabilities"), ("权益", "equity")):
            for row in bs.get(key, []):
                bs_rows.append((kind, row.get("code"), row.get("name"), _money(row.get("amount_cents", 0))))
        _section("资产负债表", ["类别", "项目/科目", "名称", "期末金额"], bs_rows)
        income = bundle["income_statement"]
        current, ytd = income.get("current", income), income.get("year_to_date", income)
        is_rows = []
        cur = {row["code"]: row for key in ("revenue", "expense") for row in current.get(key, [])}
        year = {row["code"]: row for key in ("revenue", "expense") for row in ytd.get(key, [])}
        for code in dict.fromkeys([*cur, *year]):
            row = cur.get(code) or year[code]
            kind = "收入" if row in current.get("revenue", []) or row in ytd.get("revenue", []) else "费用"
            is_rows.append((kind, code, row.get("name"), _money((cur.get(code) or {}).get("amount_cents", 0)), _money((year.get(code) or {}).get("amount_cents", 0))))
        is_rows.append(("合计", "", "净利润", _money(current.get("net_profit_cents", 0)), _money(ytd.get("net_profit_cents", 0))))
        _section("利润表", ["类别", "科目", "名称", "本月数", "本年累计"], is_rows)
        ab_rows = []
        account_table = bundle.get("account_balance_table") or {}
        for row in account_table.get("rows", []):
            opening = _row_cents(row, "opening_debit_cents", "opening_debit") - _row_cents(row, "opening_credit_cents", "opening_credit")
            ab_rows.append((row.get("code"), row.get("name"), _money(opening), _money(_row_cents(row, "period_debit_cents", "period_debit")), _money(_row_cents(row, "period_credit_cents", "period_credit")), _money(_row_cents(row, "ending_debit_cents", "ending_debit")), _money(_row_cents(row, "ending_credit_cents", "ending_credit"))))
        _section("科目余额表", ["科目编码", "科目名称", "期初余额", "本期借方", "本期贷方", "期末借方", "期末贷方"], ab_rows)
        journal_rows = []
        journal = bundle.get("chronological_ledger") or {}
        for row in journal.get("rows", []):
            journal_rows.append((row.get("date"), row.get("vno"), row.get("seq"), row.get("account_code"), row.get("account_name"), row.get("line_summary") or row.get("voucher_summary") or "", _money(_row_cents(row, "debit_cents", "debit")), _money(_row_cents(row, "credit_cents", "credit")), row.get("source") or ""))
        _section("序时簿", ["日期", "凭证号", "序号", "科目编码", "科目名称", "摘要", "借方", "贷方", "来源"], journal_rows)
        path.parent.mkdir(parents=True, exist_ok=True)
        pdf.output(str(path))

    # ------------------------------ 设置 ------------------------------
    def _page_settings(self, parent) -> None:
        self.header_status.configure(text="敏感密钥使用 Windows DPAPI 保存，界面从不回显", fg=C_TEXT_MUTED)
        canvas = tk.Canvas(parent, bg=C_BG, highlightthickness=0)
        scrollbar = ttk.Scrollbar(parent, orient="vertical", command=canvas.yview)
        inner = tk.Frame(canvas, bg=C_BG)
        window = canvas.create_window((0, 0), window=inner, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)
        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")
        inner.bind("<Configure>", lambda _e: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.bind("<Configure>", lambda e: canvas.itemconfigure(window, width=e.width))

        company = self._panel(inner, padx=16, pady=14)
        company.pack(fill="x", pady=(0, 10))
        tk.Label(company, text="公司信息", font=FONT_B, bg=C_PANEL, fg=C_TEXT).grid(row=0, column=0, columnspan=4, sticky="w", pady=(0, 8))
        conn = db.get_conn()
        row = conn.execute("SELECT * FROM company WHERE id=1").fetchone()
        conn.close()
        name_var = tk.StringVar(value=(row["name"] if row else "") or "")
        tax_var = tk.StringVar(value=(row["tax_no"] if row else "") or "")
        tk.Label(company, text="公司名称", bg=C_PANEL, font=FONT).grid(row=1, column=0, sticky="e", padx=(0, 5), pady=4)
        ttk.Entry(company, textvariable=name_var, width=34).grid(row=1, column=1, sticky="ew", padx=(0, 18), pady=4)
        tk.Label(company, text="统一社会信用代码", bg=C_PANEL, font=FONT).grid(row=1, column=2, sticky="e", padx=(0, 5), pady=4)
        ttk.Entry(company, textvariable=tax_var, width=28).grid(row=1, column=3, sticky="ew", pady=4)
        company.grid_columnconfigure(1, weight=1)
        company.grid_columnconfigure(3, weight=1)

        def save_company():
            if not name_var.get().strip():
                messagebox.showerror("无法保存", "公司名称不能为空。", parent=self.root)
                return
            try:
                updater = getattr(db, "update_company", None)
                if updater:
                    _call_supported(updater, name=name_var.get().strip(), tax_no=tax_var.get().strip())
                else:
                    conn = db.get_conn()
                    conn.execute("UPDATE company SET name=?,tax_no=? WHERE id=1", (name_var.get().strip(), tax_var.get().strip()))
                    conn.commit()
                    conn.close()
            except Exception as exc:
                messagebox.showerror("保存失败", str(exc), parent=self.root)
                return
            self.company_var.set(name_var.get().strip())
            self._clear_dirty()
            messagebox.showinfo("已保存", "公司信息已更新。", parent=self.root)

        ttk.Button(company, text="保存公司信息", command=save_company).grid(row=2, column=3, sticky="e", pady=(8, 0))
        for variable in (name_var, tax_var):
            variable.trace_add("write", lambda *_a: self._mark_dirty("公司信息"))
        self._current_save = save_company

        operations = self._panel(inner, padx=16, pady=14)
        operations.pack(fill="x", pady=10)
        tk.Label(operations, text="期间、备份与完整性", font=FONT_B, bg=C_PANEL, fg=C_TEXT).grid(row=0, column=0, columnspan=6, sticky="w", pady=(0, 8))
        period_var = tk.StringVar(value=_current_period())
        ttk.Combobox(operations, textvariable=period_var, values=self._period_values(), width=12, state="readonly").grid(row=1, column=0, padx=(0, 6), pady=4)
        ttk.Button(operations, text="结账", command=lambda: self._period_action("close", period_var.get())).grid(row=1, column=1, padx=4)
        ttk.Button(operations, text="反结账", command=lambda: self._period_action("reopen", period_var.get())).grid(row=1, column=2, padx=4)
        ttk.Button(operations, text="创建数据库备份", command=self._create_backup).grid(row=1, column=3, padx=(18, 4))
        ttk.Button(operations, text="运行完整性检查", command=self._integrity_check).grid(row=1, column=4, padx=4)
        path_text = "数据库：%s\n备份：%s\n报表：%s\n日志：%s" % (
            getattr(app_paths, "DB_PATH", getattr(db, "DB_PATH", "")),
            getattr(app_paths, "BACKUP_DIR", ""), getattr(app_paths, "REPORTS_DIR", ""), getattr(app_paths, "LOG_DIR", ""),
        )
        tk.Label(operations, text=path_text, bg=C_PANEL, fg=C_TEXT_MUTED, font=FONT_SMALL, justify="left").grid(row=2, column=0, columnspan=6, sticky="w", pady=(10, 0))

        # 开机自启开关
        autostart_var = tk.BooleanVar(value=app_paths.get_autostart_enabled())

        def toggle_autostart():
            try:
                app_paths.set_autostart_enabled(autostart_var.get())
                state = "已开启" if autostart_var.get() else "已关闭"
                messagebox.showinfo("开机自启", f"开机自启{state}。", parent=self.root)
            except Exception as exc:
                messagebox.showerror("设置失败", str(exc), parent=self.root)
                autostart_var.set(not autostart_var.get())

        tk.Checkbutton(
            operations,
            text="开机自动启动（后台运行）",
            variable=autostart_var,
            command=toggle_autostart,
            bg=C_PANEL,
            font=FONT,
        ).grid(row=3, column=0, columnspan=6, sticky="w", pady=(8, 0))

        integrations = self._panel(inner, padx=16, pady=14)
        integrations.pack(fill="x", pady=10)
        tk.Label(integrations, text="飞书审批与发票识别", font=FONT_B, bg=C_PANEL, fg=C_TEXT).grid(row=0, column=0, columnspan=4, sticky="w", pady=(0, 8))
        cfg = self._load_settings_json()
        feishu_cfg = dict(cfg.get("feishu") or {})
        fields = (
            ("App ID", "app_id"), ("审批定义 approval_code", "approval_code"),
            ("发起人 open_id", "initiator_open_id"), ("审批人 open_id", "approver_open_id"),
            ("审批节点 ID", "approval_node_id"), ("部门 ID（可选）", "department_id"),
        )
        feishu_vars = {}
        for index, (label, key) in enumerate(fields):
            row_index, col = divmod(index, 2)
            variable = tk.StringVar(value=str(feishu_cfg.get(key) or ""))
            feishu_vars[key] = variable
            tk.Label(integrations, text=label, bg=C_PANEL, font=FONT).grid(row=row_index + 1, column=col * 2, padx=(0, 5), pady=4, sticky="e")
            ttk.Entry(integrations, textvariable=variable, width=32).grid(row=row_index + 1, column=col * 2 + 1, padx=(0, 16), pady=4, sticky="ew")
            integrations.grid_columnconfigure(col * 2 + 1, weight=1)
        invoice_command_var = tk.StringVar(value=json.dumps(cfg.get("invoice_agent_command") or [], ensure_ascii=False))
        invoice_summary_var = tk.StringVar(value=str(cfg.get("invoice_summary_path") or ""))
        tk.Label(integrations, text="识别命令 JSON 数组", bg=C_PANEL, font=FONT).grid(row=4, column=0, sticky="e", padx=(0, 5), pady=4)
        ttk.Entry(integrations, textvariable=invoice_command_var).grid(row=4, column=1, columnspan=3, sticky="ew", padx=(0, 16), pady=4)
        tk.Label(integrations, text="summary.json（可选）", bg=C_PANEL, font=FONT).grid(row=5, column=0, sticky="e", padx=(0, 5), pady=4)
        ttk.Entry(integrations, textvariable=invoice_summary_var).grid(row=5, column=1, columnspan=3, sticky="ew", padx=(0, 16), pady=4)
        secret_status = tk.StringVar()

        def refresh_secret_status():
            try:
                configured = bool(security and security.get_secret("feishu.app_secret"))
                secret_status.set("App Secret：已配置" if configured else "App Secret：未配置")
            except Exception:
                secret_status.set("App Secret：状态不可用")

        def set_secret_value():
            if not security:
                messagebox.showerror("不可用", "安全凭据组件未加载。", parent=self.root)
                return
            value = simpledialog.askstring("更新 App Secret", "请输入新的飞书 App Secret（保存后不会回显）：", show="*", parent=self.root)
            if not value:
                return
            try:
                security.set_secret("feishu.app_secret", value.strip())
            except Exception as exc:
                messagebox.showerror("保存失败", str(exc), parent=self.root)
                return
            refresh_secret_status()
            messagebox.showinfo("已更新", "App Secret 已使用当前 Windows 用户的 DPAPI 加密保存。", parent=self.root)

        def delete_secret_value():
            if not security or not messagebox.askyesno("删除密钥", "确定删除已保存的飞书 App Secret？", parent=self.root):
                return
            try:
                security.delete_secret("feishu.app_secret")
            except Exception as exc:
                messagebox.showerror("删除失败", str(exc), parent=self.root)
                return
            refresh_secret_status()

        def save_integrations():
            try:
                command = json.loads(invoice_command_var.get() or "[]")
                if not isinstance(command, list) or any(not isinstance(item, str) or not item for item in command):
                    raise ValueError("识别命令必须是非空字符串组成的 JSON 数组")
                data = self._load_settings_json()
                data["feishu"] = {key: variable.get().strip() for key, variable in feishu_vars.items()}
                data["feishu"]["app_secret_ref"] = "feishu.app_secret"
                data["feishu"]["allow_self_approval"] = False
                data["invoice_agent_command"] = command
                if invoice_summary_var.get().strip():
                    data["invoice_summary_path"] = invoice_summary_var.get().strip()
                else:
                    data.pop("invoice_summary_path", None)
                self._save_settings_json(data)
                checked = feishu.validate_config()
                ok, message, _value = _result(checked)
                if not ok:
                    messagebox.showwarning("已保存但需处理", message + "\n\n请按提示修正；发起人和审批人必须使用不同 open_id。", parent=self.root)
                else:
                    messagebox.showinfo("设置已保存", message, parent=self.root)
                self._clear_dirty()
            except Exception as exc:
                messagebox.showerror("保存失败", str(exc), parent=self.root)

        row_actions = tk.Frame(integrations, bg=C_PANEL)
        row_actions.grid(row=6, column=0, columnspan=4, sticky="ew", pady=(10, 0))
        tk.Label(row_actions, textvariable=secret_status, bg=C_PANEL, fg=C_TEXT_MUTED, font=FONT_SMALL).pack(side="left")
        ttk.Button(row_actions, text="更新密钥", command=set_secret_value).pack(side="left", padx=6)
        ttk.Button(row_actions, text="删除密钥", command=delete_secret_value).pack(side="left")
        ttk.Button(row_actions, text="保存集成设置", command=save_integrations).pack(side="right")
        refresh_secret_status()
        for variable in (*feishu_vars.values(), invoice_command_var, invoice_summary_var):
            variable.trace_add("write", lambda *_a: self._mark_dirty("集成设置"))

        issues_panel = self._panel(inner, padx=16, pady=14)
        issues_panel.pack(fill="x", pady=(10, 20))
        tk.Label(issues_panel, text="数据异常", font=FONT_B, bg=C_PANEL, fg=C_TEXT).pack(side="left")
        tk.Label(issues_panel, text=f"当前阻断项：{len(self._issues)}", font=FONT, bg=C_PANEL, fg=C_DANGER if self._issues else C_OK).pack(side="left", padx=16)
        ttk.Button(issues_panel, text="查看并处理", command=self._open_issue_dialog).pack(side="right")

    def _settings_path(self) -> Path:
        directory = Path(getattr(app_paths, "CONFIG_DIR", ROOT / "config"))
        directory.mkdir(parents=True, exist_ok=True)
        return directory / "settings.json"

    def _load_settings_json(self) -> dict:
        path = self._settings_path()
        if not path.exists():
            return {}
        try:
            value = json.loads(path.read_text(encoding="utf-8"))
            return value if isinstance(value, dict) else {}
        except (OSError, ValueError):
            return {}

    def _save_settings_json(self, value: dict) -> None:
        path = self._settings_path()
        temporary = path.with_suffix(".tmp")
        temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
        os.replace(temporary, path)

    def _period_action(self, action: str, period: str) -> None:
        try:
            period = _validate_period(period)
        except ValueError as exc:
            messagebox.showerror("期间无效", str(exc), parent=self.root)
            return
        dialog = WorkflowDialog(self.root, "结账" if action == "close" else "反结账", self.default_actor)
        if not dialog.result:
            return
        try:
            if action == "close":
                result = _call_supported(ledger.close_period, period=period, actor=dialog.result["actor"], reason=dialog.result["reason"] or None)
            else:
                if not dialog.result["reason"]:
                    raise ValueError("反结账必须填写原因")
                result = _call_supported(ledger.reopen_period, period=period, actor=dialog.result["actor"], reason=dialog.result["reason"])
            ok, message, _value = _result(result)
            if not ok:
                raise ValueError(message)
        except Exception as exc:
            messagebox.showerror("操作失败", str(exc), parent=self.root)
            return
        messagebox.showinfo("期间操作", message or "操作成功。", parent=self.root)

    def _create_backup(self) -> None:
        initial_dir = str(getattr(app_paths, "BACKUP_DIR", ROOT / "backups"))
        path = filedialog.asksaveasfilename(parent=self.root, title="创建数据库备份", initialdir=initial_dir, initialfile=f"finance_{dt.datetime.now():%Y%m%d_%H%M%S}.db", defaultextension=".db", filetypes=(("SQLite 数据库", "*.db"),))
        if not path:
            return
        functions = ("backup_database", "create_backup", "backup_db")
        function = next((getattr(db, name, None) for name in functions if callable(getattr(db, name, None))), None)
        if not function:
            messagebox.showerror("备份不可用", "数据库层未提供 SQLite Backup API。", parent=self.root)
            return
        try:
            signature = inspect.signature(function)
            kwargs = {"destination": path, "destination_path": path, "target_path": path, "path": path}
            result = function(**{key: value for key, value in kwargs.items() if key in signature.parameters})
            ok, message, _value = _result(result)
            if not ok:
                raise ValueError(message)
        except Exception as exc:
            messagebox.showerror("备份失败", str(exc), parent=self.root)
            return
        messagebox.showinfo("备份完成", message or f"备份已保存到：\n{path}", parent=self.root)

    def _integrity_check(self) -> None:
        function = next((getattr(db, name, None) for name in ("run_integrity_checks", "check_integrity", "integrity_check") if callable(getattr(db, name, None))), None)
        try:
            if function:
                result = function()
                ok, message, data = _result(result)
                detail = json.dumps(data, ensure_ascii=False, indent=2) if data not in (None, "") else message
                if not ok:
                    raise ValueError(detail or message)
            else:
                conn = db.get_conn()
                integrity = conn.execute("PRAGMA integrity_check").fetchone()[0]
                foreign = conn.execute("PRAGMA foreign_key_check").fetchall()
                conn.close()
                if integrity != "ok" or foreign:
                    raise ValueError(f"完整性：{integrity}；外键异常：{len(foreign)}")
                detail = "SQLite 完整性检查通过，未发现外键异常。"
        except Exception as exc:
            messagebox.showerror("检查发现问题", str(exc), parent=self.root)
            return
        self._refresh_issue_state()
        messagebox.showinfo("完整性检查", detail, parent=self.root)

    # ------------------------------ dialogs ------------------------------
    def _open_issue_dialog(self) -> None:
        self._refresh_issue_state()
        if not self._issues:
            messagebox.showinfo("数据异常", "当前没有阻断性数据异常。", parent=self.root)
            return
        IssueResolutionDialog(self, self._issues)

    def _add_partner_dialog(self, on_saved=None) -> None:
        dialog = PartnerDialog(self.root)
        if not dialog.result:
            return
        try:
            arap.add_partner(**dialog.result)
        except Exception as exc:
            messagebox.showerror("保存失败", str(exc), parent=self.root)
            return
        if callable(on_saved):
            on_saved()

    def _add_employee_dialog(self) -> None:
        dialog = EmployeeDialog(self.root)
        if not dialog.result:
            return
        try:
            expense_mod.add_employee(**dialog.result)
        except Exception as exc:
            messagebox.showerror("保存失败", str(exc), parent=self.root)
            return
        messagebox.showinfo("已保存", "员工档案已保存。", parent=self.root)

    def _add_expense_dialog(self, on_saved=None) -> None:
        employees = expense_mod.list_employees()
        if not employees:
            messagebox.showwarning("缺少员工", "请先新增员工。", parent=self.root)
            return
        dialog = ExpenseDialog(self.root, employees, self._accounts(category="损益", direction="借"))
        if not dialog.result:
            return
        try:
            result = _call_supported(expense_mod.create_expense, **dialog.result)
            ok, message, _value = _result(result)
            if not ok:
                raise ValueError(message)
        except Exception as exc:
            messagebox.showerror("提交失败", str(exc), parent=self.root)
            return
        messagebox.showinfo("报销单已提交", message, parent=self.root)
        if callable(on_saved):
            on_saved()


class LineDialog:
    def __init__(self, parent, accounts: list[dict], default_summary="", initial=None):
        self.result = None
        window = tk.Toplevel(parent)
        window.configure(bg=C_BG)
        window.title("分录")
        window.transient(parent)
        window.grab_set()
        window.resizable(False, False)
        values = [FinanceApp._account_display(account) for account in accounts]
        account_var = tk.StringVar()
        if initial:
            account_var.set(next((value for value in values if value.startswith(str(initial.get("account_code")) + " ")), values[0] if values else ""))
        else:
            account_var.set(values[0] if values else "")
        summary_var = tk.StringVar(value=(initial or {}).get("summary", default_summary))
        debit_var = tk.StringVar(value=_money((initial or {}).get("debit_cents", 0)).replace(",", ""))
        credit_var = tk.StringVar(value=_money((initial or {}).get("credit_cents", 0)).replace(",", ""))
        aux_var = tk.StringVar(value=(initial or {}).get("aux_type") or "无")
        aux_id_var = tk.StringVar(value=str((initial or {}).get("aux_id") or ""))
        fields = (("科目", ttk.Combobox(window, textvariable=account_var, values=values, width=30, state="readonly")), ("摘要", ttk.Entry(window, textvariable=summary_var, width=33)), ("借方金额", ttk.Entry(window, textvariable=debit_var, width=18)), ("贷方金额", ttk.Entry(window, textvariable=credit_var, width=18)), ("辅助类型", ttk.Combobox(window, textvariable=aux_var, values=("无", "客户", "供应商", "员工"), width=16, state="readonly")), ("辅助 ID", ttk.Entry(window, textvariable=aux_id_var, width=18)))
        for row, (label, widget) in enumerate(fields):
            tk.Label(window, text=label, font=FONT).grid(row=row, column=0, padx=10, pady=6, sticky="e")
            widget.grid(row=row, column=1, padx=10, pady=6, sticky="w")

        def accept():
            try:
                debit = _parse_cents(debit_var.get())
                credit = _parse_cents(credit_var.get())
                if bool(debit) == bool(credit):
                    raise ValueError("借方和贷方必须且只能填写一方")
                if aux_var.get() != "无" and not aux_id_var.get().strip().isdigit():
                    raise ValueError("选择辅助核算后必须填写有效的辅助 ID")
                self.result = {"account_code": account_var.get().split()[0], "summary": summary_var.get().strip(), "debit_cents": debit, "credit_cents": credit, "aux_type": aux_var.get(), "aux_id": int(aux_id_var.get()) if aux_id_var.get().strip() else None}
            except Exception as exc:
                messagebox.showerror("分录无效", str(exc), parent=window)
                return
            window.destroy()

        ttk.Button(window, text="取消", command=window.destroy).grid(row=len(fields), column=0, padx=10, pady=12, sticky="e")
        ttk.Button(window, text="确定", command=accept).grid(row=len(fields), column=1, padx=10, pady=12, sticky="e")
        window.bind("<Return>", lambda _e: accept())
        window.bind("<Escape>", lambda _e: window.destroy())
        window.wait_window()


class WorkflowDialog:
    def __init__(self, parent, action: str, default_actor: str):
        self.result = None
        window = tk.Toplevel(parent)
        window.configure(bg=C_BG)
        window.title(action)
        window.transient(parent)
        window.grab_set()
        window.resizable(False, False)
        actor_var = tk.StringVar(value=default_actor)
        reason_var = tk.StringVar()
        date_var = tk.StringVar(value=_today())
        tk.Label(window, text="操作人", font=FONT).grid(row=0, column=0, padx=10, pady=7, sticky="e")
        ttk.Entry(window, textvariable=actor_var, width=30).grid(row=0, column=1, padx=10, pady=7)
        row = 1
        if action == "冲销":
            tk.Label(window, text="冲销日期", font=FONT).grid(row=row, column=0, padx=10, pady=7, sticky="e")
            ttk.Entry(window, textvariable=date_var, width=30).grid(row=row, column=1, padx=10, pady=7)
            row += 1
        tk.Label(window, text="原因/兼任说明", font=FONT).grid(row=row, column=0, padx=10, pady=7, sticky="e")
        ttk.Entry(window, textvariable=reason_var, width=42).grid(row=row, column=1, padx=10, pady=7)
        tk.Label(window, text="审核人与制单人相同时必须填写兼任说明；反结账和冲销必须填写原因。", font=FONT_SMALL, fg=C_TEXT_MUTED, wraplength=420, justify="left").grid(row=row + 1, column=0, columnspan=2, padx=12, pady=(0, 8), sticky="w")

        def accept():
            try:
                if not actor_var.get().strip():
                    raise ValueError("操作人不能为空")
                if action in ("冲销", "反结账", "删除草稿") and not reason_var.get().strip():
                    raise ValueError("必须填写原因")
                date_value = _validate_date(date_var.get()) if action == "冲销" else ""
                self.result = {"actor": actor_var.get().strip(), "reason": reason_var.get().strip(), "date": date_value}
            except ValueError as exc:
                messagebox.showerror("信息不完整", str(exc), parent=window)
                return
            window.destroy()

        ttk.Button(window, text="取消", command=window.destroy).grid(row=row + 2, column=0, padx=10, pady=12, sticky="e")
        ttk.Button(window, text="确认", command=accept).grid(row=row + 2, column=1, padx=10, pady=12, sticky="e")
        window.wait_window()


class PartnerDialog:
    def __init__(self, parent):
        self.result = None
        window = tk.Toplevel(parent)
        window.configure(bg=C_BG)
        window.title("新增往来方")
        window.transient(parent)
        window.grab_set()
        variables = {"ptype": tk.StringVar(value="客户"), "name": tk.StringVar(), "tax_no": tk.StringVar(), "contact": tk.StringVar(), "phone": tk.StringVar()}
        fields = (("类型", ttk.Combobox(window, textvariable=variables["ptype"], values=("客户", "供应商"), state="readonly", width=23)), ("名称", ttk.Entry(window, textvariable=variables["name"], width=26)), ("税号", ttk.Entry(window, textvariable=variables["tax_no"], width=26)), ("联系人", ttk.Entry(window, textvariable=variables["contact"], width=26)), ("电话", ttk.Entry(window, textvariable=variables["phone"], width=26)))
        for row, (label, widget) in enumerate(fields):
            tk.Label(window, text=label, font=FONT).grid(row=row, column=0, padx=10, pady=6, sticky="e")
            widget.grid(row=row, column=1, padx=10, pady=6)

        def accept():
            if not variables["name"].get().strip():
                messagebox.showerror("信息不完整", "名称不能为空。", parent=window)
                return
            self.result = {key: variable.get().strip() for key, variable in variables.items()}
            window.destroy()

        ttk.Button(window, text="取消", command=window.destroy).grid(row=5, column=0, padx=10, pady=12, sticky="e")
        ttk.Button(window, text="保存", command=accept).grid(row=5, column=1, padx=10, pady=12, sticky="e")
        window.wait_window()


class EmployeeDialog:
    def __init__(self, parent):
        self.result = None
        window = tk.Toplevel(parent)
        window.configure(bg=C_BG)
        window.title("新增员工")
        window.transient(parent)
        window.grab_set()
        variables = {"name": tk.StringVar(), "dept": tk.StringVar(), "phone": tk.StringVar()}
        for row, (key, label) in enumerate((("name", "姓名"), ("dept", "部门"), ("phone", "电话"))):
            tk.Label(window, text=label, font=FONT).grid(row=row, column=0, padx=10, pady=7, sticky="e")
            ttk.Entry(window, textvariable=variables[key], width=28).grid(row=row, column=1, padx=10, pady=7)

        def accept():
            if not variables["name"].get().strip():
                messagebox.showerror("信息不完整", "姓名不能为空。", parent=window)
                return
            self.result = {key: variable.get().strip() for key, variable in variables.items()}
            window.destroy()

        ttk.Button(window, text="取消", command=window.destroy).grid(row=3, column=0, padx=10, pady=12, sticky="e")
        ttk.Button(window, text="保存", command=accept).grid(row=3, column=1, padx=10, pady=12, sticky="e")
        window.wait_window()


class ExpenseDialog:
    def __init__(self, parent, employees: list[dict], accounts: list[dict]):
        self.result = None
        window = tk.Toplevel(parent)
        window.configure(bg=C_BG)
        window.title("新增报销单")
        window.geometry("720x500")
        window.transient(parent)
        window.grab_set()
        emp_values = [f"{row['id']}|{row['name']}" for row in employees]
        account_values = [FinanceApp._account_display(row) for row in accounts]
        employee_var = tk.StringVar(value=emp_values[0])
        date_var = tk.StringVar(value=_today())
        summary_var = tk.StringVar()
        tk.Label(window, text="员工", font=FONT).grid(row=0, column=0, padx=8, pady=8, sticky="e")
        ttk.Combobox(window, textvariable=employee_var, values=emp_values, state="readonly", width=22).grid(row=0, column=1, padx=8, pady=8, sticky="w")
        tk.Label(window, text="日期", font=FONT).grid(row=0, column=2, padx=8, pady=8, sticky="e")
        ttk.Entry(window, textvariable=date_var, width=14).grid(row=0, column=3, padx=8, pady=8, sticky="w")
        tk.Label(window, text="事由", font=FONT).grid(row=1, column=0, padx=8, pady=8, sticky="e")
        ttk.Entry(window, textvariable=summary_var).grid(row=1, column=1, columnspan=3, padx=8, pady=8, sticky="ew")
        tree = ttk.Treeview(window, columns=("account", "summary", "amount", "invoice"), show="headings", height=10)
        for key, label, width in (("account", "费用科目", 190), ("summary", "摘要", 220), ("amount", "金额", 100), ("invoice", "发票号", 130)):
            tree.heading(key, text=label)
            tree.column(key, width=width)
        tree.grid(row=2, column=0, columnspan=4, padx=8, pady=8, sticky="nsew")
        window.grid_rowconfigure(2, weight=1)
        window.grid_columnconfigure(3, weight=1)
        items: list[dict] = []

        def add_item():
            dialog = ExpenseItemDialog(window, account_values)
            if dialog.result:
                items.append(dialog.result)
                tree.insert("", "end", iid=str(len(items) - 1), values=(dialog.result["account_code"], dialog.result["summary"], _money(dialog.result["amount_cents"]), dialog.result["invoice_ref"]))

        def remove_item():
            selection = tree.selection()
            if not selection:
                return
            index = int(selection[0])
            del items[index]
            tree.delete(*tree.get_children())
            for position, item in enumerate(items):
                tree.insert("", "end", iid=str(position), values=(item["account_code"], item["summary"], _money(item["amount_cents"]), item["invoice_ref"]))

        buttons = tk.Frame(window, bg=C_BG)
        buttons.grid(row=3, column=0, columnspan=4, sticky="ew", padx=8)
        ttk.Button(buttons, text="增加明细", command=add_item).pack(side="left")
        ttk.Button(buttons, text="删除明细", command=remove_item).pack(side="left", padx=6)

        def accept():
            try:
                date_value = _validate_date(date_var.get())
                if not items:
                    raise ValueError("至少需要一条报销明细")
                if not summary_var.get().strip():
                    raise ValueError("报销事由不能为空")
                normalized = [{**item, "amount": item["amount_cents"] / 100.0} for item in items]
                self.result = {"emp_id": int(employee_var.get().split("|", 1)[0]), "employee_id": int(employee_var.get().split("|", 1)[0]), "date_str": date_value, "items": normalized, "summary": summary_var.get().strip(), "approver": ""}
            except ValueError as exc:
                messagebox.showerror("信息不完整", str(exc), parent=window)
                return
            window.destroy()

        ttk.Button(buttons, text="取消", command=window.destroy).pack(side="right")
        ttk.Button(buttons, text="提交审批", command=accept).pack(side="right", padx=6)
        window.wait_window()


class ExpenseItemDialog:
    def __init__(self, parent, account_values: list[str]):
        self.result = None
        window = tk.Toplevel(parent)
        window.configure(bg=C_BG)
        window.title("报销明细")
        window.transient(parent)
        window.grab_set()
        account_var = tk.StringVar(value=account_values[0] if account_values else "")
        summary_var, amount_var, invoice_var = tk.StringVar(), tk.StringVar(), tk.StringVar()
        fields = (("费用科目", ttk.Combobox(window, textvariable=account_var, values=account_values, width=28, state="readonly")), ("摘要", ttk.Entry(window, textvariable=summary_var, width=31)), ("金额（元）", ttk.Entry(window, textvariable=amount_var, width=18)), ("发票号", ttk.Entry(window, textvariable=invoice_var, width=31)))
        for row, (label, widget) in enumerate(fields):
            tk.Label(window, text=label, font=FONT).grid(row=row, column=0, padx=10, pady=7, sticky="e")
            widget.grid(row=row, column=1, padx=10, pady=7, sticky="w")

        def accept():
            try:
                if not account_var.get():
                    raise ValueError("请选择费用科目")
                amount = _parse_cents(amount_var.get(), positive=True)
                self.result = {"account_code": account_var.get().split()[0], "summary": summary_var.get().strip(), "amount_cents": amount, "invoice_ref": invoice_var.get().strip()}
            except ValueError as exc:
                messagebox.showerror("明细无效", str(exc), parent=window)
                return
            window.destroy()

        ttk.Button(window, text="取消", command=window.destroy).grid(row=4, column=0, padx=10, pady=12, sticky="e")
        ttk.Button(window, text="确定", command=accept).grid(row=4, column=1, padx=10, pady=12, sticky="e")
        window.wait_window()


class InvoiceConfirmDialog:
    def __init__(self, parent, accounts: list[dict], suppliers: list[dict], default_actor: str):
        self.result = None
        window = tk.Toplevel(parent)
        window.configure(bg=C_BG)
        window.title("确认发票入账信息")
        window.transient(parent)
        window.grab_set()
        account_values = [FinanceApp._account_display(row) for row in accounts]
        supplier_values = [f"{row['id']}|{row['name']}" for row in suppliers]
        account_var = tk.StringVar(value=account_values[0] if account_values else "")
        supplier_var = tk.StringVar(value=supplier_values[0] if supplier_values else "")
        payment_var = tk.StringVar(value="未付款")
        actor_var = tk.StringVar(value=default_actor)
        tax_var = tk.BooleanVar(value=False)
        fields = (("费用科目", ttk.Combobox(window, textvariable=account_var, values=account_values, width=30, state="readonly")), ("供应商", ttk.Combobox(window, textvariable=supplier_var, values=supplier_values, width=30, state="readonly")), ("付款状态", ttk.Combobox(window, textvariable=payment_var, values=("未付款", "已付款"), width=18, state="readonly")), ("制单人", ttk.Entry(window, textvariable=actor_var, width=24)))
        for row, (label, widget) in enumerate(fields):
            tk.Label(window, text=label, font=FONT).grid(row=row, column=0, padx=10, pady=7, sticky="e")
            widget.grid(row=row, column=1, padx=10, pady=7, sticky="w")
        ttk.Checkbutton(window, text="税额可抵扣（拆分进项税）", variable=tax_var).grid(row=4, column=1, padx=10, pady=5, sticky="w")
        tk.Label(window, text="默认“未付款”，贷记应付账款；仅在已实际支付时选择“已付款”。", font=FONT_SMALL, fg=C_TEXT_MUTED, wraplength=400, justify="left").grid(row=5, column=0, columnspan=2, padx=10, pady=6, sticky="w")

        def accept():
            if not supplier_var.get():
                messagebox.showerror("信息不完整", "请先在往来管理中新增并选择供应商。", parent=window)
                return
            if not account_var.get() or not actor_var.get().strip():
                messagebox.showerror("信息不完整", "费用科目和制单人不能为空。", parent=window)
                return
            self.result = {"expense_account": account_var.get().split()[0], "partner_id": int(supplier_var.get().split("|", 1)[0]), "payment_status": payment_var.get(), "deductible_tax": tax_var.get(), "actor": actor_var.get().strip()}
            window.destroy()

        ttk.Button(window, text="取消", command=window.destroy).grid(row=6, column=0, padx=10, pady=12, sticky="e")
        ttk.Button(window, text="确认", command=accept).grid(row=6, column=1, padx=10, pady=12, sticky="e")
        window.wait_window()


class IssueResolutionDialog:
    def __init__(self, app: FinanceApp, issues: list[dict]):
        self.app = app
        window = tk.Toplevel(app.root)
        window.configure(bg=C_BG)
        window.title("处理阻断性数据异常")
        window.geometry("840x520")
        window.transient(app.root)
        window.grab_set()
        tk.Label(window, text="正式报表已暂停。请逐项确认正确科目，并记录修复人和原因。", font=FONT_B, fg=C_DANGER).pack(anchor="w", padx=14, pady=(14, 8))
        panel = tk.Frame(window, bg=C_BG)
        panel.pack(fill="both", expand=True, padx=14)
        tree = ttk.Treeview(panel, columns=("id", "code", "entity", "detail"), show="headings", height=10)
        for key, label, width in (("id", "ID", 55), ("code", "异常代码", 175), ("entity", "关联对象", 170), ("detail", "说明", 390)):
            tree.heading(key, text=label)
            tree.column(key, width=width)
        ybar = ttk.Scrollbar(panel, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=ybar.set)
        tree.pack(side="left", fill="both", expand=True)
        ybar.pack(side="right", fill="y")
        issue_map = {}
        for issue in issues:
            iid = str(issue.get("id"))
            issue_map[iid] = issue
            entity = issue.get("voucher_no") or issue.get("entity_key") or issue.get("entity_id") or issue.get("ref_id") or ""
            detail = issue.get("message") or issue.get("description") or issue.get("issue_detail") or (f"共 {issue.get('count')} 条" if issue.get("count") else "需要人工确认")
            tree.insert("", "end", iid=iid, values=(issue.get("id"), issue.get("issue_code") or issue.get("code"), entity, detail))
        form = tk.Frame(window, bg=C_BG)
        form.pack(fill="x", padx=14, pady=12)
        accounts = app._accounts()
        account_values = [FinanceApp._account_display(row) for row in accounts]
        account_var = tk.StringVar(value=account_values[0] if account_values else "")
        actor_var = tk.StringVar(value=app.default_actor)
        reason_var = tk.StringVar()
        for index, (label, widget) in enumerate((("正确科目", ttk.Combobox(form, textvariable=account_var, values=account_values, width=32, state="readonly")), ("修复人", ttk.Entry(form, textvariable=actor_var, width=24)), ("修复原因", ttk.Entry(form, textvariable=reason_var, width=46)))):
            tk.Label(form, text=label, font=FONT).grid(row=index, column=0, padx=(0, 6), pady=5, sticky="e")
            widget.grid(row=index, column=1, padx=(0, 12), pady=5, sticky="ew")
        form.grid_columnconfigure(1, weight=1)

        def resolve():
            selection = tree.selection()
            if not selection:
                messagebox.showwarning("请选择异常", "请先选择一项异常。", parent=window)
                return
            if not actor_var.get().strip() or not reason_var.get().strip():
                messagebox.showerror("信息不完整", "修复人和修复原因均为必填项。", parent=window)
                return
            resolver = getattr(db, "resolve_data_issue", None)
            if not resolver:
                messagebox.showerror("修复不可用", "数据库迁移尚未提供异常修复服务。", parent=window)
                return
            issue = issue_map[selection[0]]
            try:
                result = resolver(issue["id"], account_var.get().split()[0], actor_var.get().strip(), reason_var.get().strip())
                ok, message, _value = _result(result)
                if not ok:
                    raise ValueError(message)
            except Exception as exc:
                messagebox.showerror("修复失败", str(exc), parent=window)
                return
            tree.delete(selection[0])
            issue_map.pop(selection[0], None)
            reason_var.set("")
            app._refresh_issue_state()
            if not issue_map:
                messagebox.showinfo("异常已处理", "所有阻断性异常均已处理，正式报表已恢复。", parent=window)
                window.destroy()
                app._show_page(app._active_nav, force=True)

        buttons = tk.Frame(window, bg=C_BG)
        buttons.pack(fill="x", padx=14, pady=(0, 14))
        ttk.Button(buttons, text="关闭", command=window.destroy).pack(side="right")
        tk.Button(buttons, text="确认修复选中项", font=FONT_B, bg=C_DANGER, fg="white", relief="flat", padx=12, pady=6, command=resolve).pack(side="right", padx=8)
        if tree.get_children():
            tree.selection_set(tree.get_children()[0])
        window.wait_window()


def main() -> None:
    _enable_dpi_awareness()
    root = tk.Tk()
    FinanceApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()

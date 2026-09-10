# -*- coding: utf-8 -*-
"""报表导出服务：Excel / PDF。

注：GUI 报表页使用 gui.main 内的 ``_export_report_*`` 函数（基于已生成的
报表包）。本模块提供更细粒度的独立导出（科目余额表 / 凭证明细 / 发票明细），
同样基于真实 schema（整数分 ``debit_cents`` / ``credit_cents``、凭证号 ``vno``、
发票 ``seller`` / ``total_cents`` 等）。
"""
from __future__ import annotations

import sqlite3
from datetime import date as _date
from pathlib import Path
from typing import Any

from modules.integration import IntegrationResult, IntegrationStatus

try:
    import logging_setup  # type: ignore

    LOG = logging_setup.get_logger("export_service")
except Exception:  # pragma: no cover - fallback when not bundled
    import logging as _logging

    LOG = _logging.getLogger("export_service")


def _check_deps() -> IntegrationResult:
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        return IntegrationResult(IntegrationStatus.CONFIG_ERROR, "缺少 openpyxl，请运行: pip install openpyxl")
    try:
        import fpdf  # noqa: F401
    except ImportError:
        return IntegrationResult(IntegrationStatus.CONFIG_ERROR, "缺少 fpdf2，请运行: pip install fpdf2")
    return IntegrationResult(IntegrationStatus.SUCCESS, "ok")


def _get_conn() -> sqlite3.Connection | None:
    try:
        import db
        return db.get_conn()
    except Exception as exc:
        LOG.error("获取数据库连接失败: %s", exc)
        return None


# ---------- 科目余额表 ----------

def _current_period() -> str:
    return _date.today().strftime("%Y-%m")


def export_balance_sheet(path: str | Path, fmt: str = "xlsx", period: str | None = None) -> IntegrationResult:
    """导出科目余额表（期末余额，单位：元）。"""
    dep = _check_deps()
    if not dep.ok:
        return dep

    path = Path(path)
    conn = _get_conn()
    if conn is None:
        return IntegrationResult(IntegrationStatus.RETRYABLE_ERROR, "数据库连接失败")

    try:
        end_period = period or _current_period()
        rows = conn.execute(
            """
            SELECT a.code, a.name, a.category,
                   COALESCE(SUM(CASE WHEN v.status IN ('posted','已记账') AND v.period<=?
                                     THEN vl.debit_cents ELSE 0 END),0) d,
                   COALESCE(SUM(CASE WHEN v.status IN ('posted','已记账') AND v.period<=?
                                     THEN vl.credit_cents ELSE 0 END),0) c
            FROM account a
            LEFT JOIN voucher_line vl ON vl.account_code = a.code
            LEFT JOIN voucher v ON v.id = vl.voucher_id
            GROUP BY a.code, a.name, a.category
            ORDER BY a.code
            """,
            (end_period, end_period),
        ).fetchall()
    except Exception as exc:
        LOG.error("查询科目余额失败: %s", exc)
        conn.close()
        return IntegrationResult(IntegrationStatus.RETRYABLE_ERROR, f"查询失败: {exc}")
    finally:
        conn.close()

    headers = ["科目编码", "科目名称", "类别", "期末余额(元)"]
    data = []
    for r in rows:
        net = int(r["d"] or 0) - int(r["c"] or 0)
        if net == 0 and r["category"] not in ("资产", "负债", "权益", "成本"):
            continue
        data.append((r["code"], r["name"], r["category"], f"{net / 100:.2f}"))

    if fmt == "xlsx":
        return _write_xlsx(path, "科目余额表", headers, data)
    if fmt == "pdf":
        return _write_pdf(path, "科目余额表", headers, data)
    return IntegrationResult(IntegrationStatus.BUSINESS_REJECTED, f"不支持的格式: {fmt}")


# ---------- 凭证明细 ----------

def export_vouchers(path: str | Path, fmt: str = "xlsx", period: str | None = None) -> IntegrationResult:
    """导出凭证明细。"""
    dep = _check_deps()
    if not dep.ok:
        return dep

    path = Path(path)
    conn = _get_conn()
    if conn is None:
        return IntegrationResult(IntegrationStatus.RETRYABLE_ERROR, "数据库连接失败")

    try:
        sql = """
            SELECT v.id, v.vno, v.period, v.date, v.summary, v.status,
                   GROUP_CONCAT(vl.account_code || ' ' || COALESCE(a.name,'') || ' 借' ||
                                 (vl.debit_cents/100.0) || ' 贷' || (vl.credit_cents/100.0), ' | ') AS lines
            FROM voucher v
            LEFT JOIN voucher_line vl ON vl.voucher_id = v.id
            LEFT JOIN account a ON a.code = vl.account_code
        """
        args: list[Any] = []
        if period:
            sql += " WHERE v.period = ?"
            args.append(period)
        sql += " GROUP BY v.id ORDER BY v.date DESC, v.vno DESC"
        rows = conn.execute(sql, args).fetchall()
    except Exception as exc:
        LOG.error("查询凭证失败: %s", exc)
        conn.close()
        return IntegrationResult(IntegrationStatus.RETRYABLE_ERROR, f"查询失败: {exc}")
    finally:
        conn.close()

    headers = ["凭证号", "期间", "日期", "摘要", "状态", "分录明细"]
    data = [
        (r["vno"], r["period"], r["date"], r["summary"], r["status"], r["lines"] or "")
        for r in rows
    ]

    if fmt == "xlsx":
        return _write_xlsx(path, "凭证明细", headers, data)
    if fmt == "pdf":
        return _write_pdf(path, "凭证明细", headers, data)
    return IntegrationResult(IntegrationStatus.BUSINESS_REJECTED, f"不支持的格式: {fmt}")


# ---------- 发票明细 ----------

def export_invoices(path: str | Path, fmt: str = "xlsx", status: str | None = None) -> IntegrationResult:
    """导出发票明细。"""
    dep = _check_deps()
    if not dep.ok:
        return dep

    path = Path(path)
    conn = _get_conn()
    if conn is None:
        return IntegrationResult(IntegrationStatus.RETRYABLE_ERROR, "数据库连接失败")

    try:
        sql = (
            "SELECT invoice_no, date, seller, buyer, "
            "total_cents, tax_cents, status, created_at FROM invoice_doc"
        )
        args: list[Any] = []
        if status:
            sql += " WHERE status = ?"
            args.append(status)
        sql += " ORDER BY date DESC, created_at DESC"
        rows = conn.execute(sql, args).fetchall()
    except Exception as exc:
        LOG.error("查询发票失败: %s", exc)
        conn.close()
        return IntegrationResult(IntegrationStatus.RETRYABLE_ERROR, f"查询失败: {exc}")
    finally:
        conn.close()

    headers = ["发票号码", "开票日期", "销售方", "购买方", "价税合计(元)", "税额(元)", "状态", "入库时间"]
    data = [
        (
            r["invoice_no"],
            r["date"],
            r["seller"],
            r["buyer"],
            f"{(int(r['total_cents'] or 0)) / 100:.2f}",
            f"{(int(r['tax_cents'] or 0)) / 100:.2f}",
            r["status"],
            r["created_at"],
        )
        for r in rows
    ]

    if fmt == "xlsx":
        return _write_xlsx(path, "发票明细", headers, data)
    if fmt == "pdf":
        return _write_pdf(path, "发票明细", headers, data)
    return IntegrationResult(IntegrationStatus.BUSINESS_REJECTED, f"不支持的格式: {fmt}")


# ---------- 底层写入 ----------

def _write_xlsx(path: Path, sheet_name: str, headers: list[str], data: list[tuple]) -> IntegrationResult:
    try:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name
        ws.append(headers)
        for row in data:
            ws.append(row)
        # Auto width (best effort)
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    max_length = max(max_length, len(str(cell.value or "")))
                except Exception:
                    pass
            ws.column_dimensions[column_letter].width = min(max_length + 2, 60)
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(path))
        return IntegrationResult(IntegrationStatus.SUCCESS, f"已导出: {path}")
    except Exception as exc:
        LOG.error("写入 Excel 失败: %s", exc)
        return IntegrationResult(IntegrationStatus.RETRYABLE_ERROR, f"Excel 导出失败: {exc}")


_CJK_FONT_CANDIDATES = (
    r"C:\Windows\Fonts\simhei.ttf",
    r"C:\Windows\Fonts\Deng.ttf",
    r"C:\Windows\Fonts\simkai.ttf",
    r"C:\Windows\Fonts\msyh.ttc",
    r"C:\Windows\Fonts\simsun.ttc",
    "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
)


def _find_cjk_font() -> str | None:
    """找到可用的中文字体文件；找不到返回 None。"""
    for candidate in _CJK_FONT_CANDIDATES:
        if Path(candidate).is_file():
            return candidate
    return None


def _write_pdf(path: Path, title: str, headers: list[str], data: list[tuple]) -> IntegrationResult:
    try:
        from fpdf import FPDF
        pdf = FPDF()
        pdf.add_page()
        pdf.set_auto_page_break(auto=True, margin=15)

        # 中文字体：优先用系统自带 TTF（Helvetica 无法渲染中文）。
        font_path = _find_cjk_font()
        if font_path:
            pdf.add_font("cj", "", font_path)
            base_font = "cj"
        else:
            base_font = "Helvetica"  # 退化为拉丁字体（仅英文可显示）

        pdf.set_font(base_font, "", 14)
        pdf.cell(0, 10, title, ln=True, align="L")
        pdf.ln(2)

        pdf.set_font(base_font, "", 9)
        col_widths = _pdf_col_widths(headers)
        for h, w in zip(headers, col_widths):
            pdf.cell(w, 7, str(h), border=1, align="C")
        pdf.ln()

        pdf.set_font(base_font, "", 9)
        for row in data:
            for value, w in zip(row, col_widths):
                text = str(value) if value is not None else ""
                pdf.cell(w, 7, text, border=1, align="L")
            pdf.ln()

        path.parent.mkdir(parents=True, exist_ok=True)
        pdf.output(str(path))
        return IntegrationResult(IntegrationStatus.SUCCESS, f"已导出: {path}")
    except Exception as exc:
        LOG.error("写入 PDF 失败: %s", exc)
        return IntegrationResult(IntegrationStatus.RETRYABLE_ERROR, f"PDF 导出失败: {exc}")


def _pdf_col_widths(headers: list[str]) -> list[float]:
    """简易等宽列，合计 190mm (A4 宽 210mm - 边距)。"""
    n = len(headers)
    base = 190 / n
    return [base] * n

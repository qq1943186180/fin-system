# -*- coding: utf-8 -*-
"""账簿（科目余额表 + 序时簿）导出回归：CSV / Excel / PDF 三路。

锁住曾经的坑：openpyxl/fpdf2 依赖缺失、PDF CJK 字体。
"""
from __future__ import annotations

import importlib

import pytest


@pytest.fixture()
def finance_modules(tmp_path, monkeypatch):
    monkeypatch.setenv("FINANCE_SYSTEM_HOME", str(tmp_path / "home"))
    monkeypatch.setenv("FINANCE_SYSTEM_TEST_MODE", "1")

    import app_paths
    import db
    from modules import ledger, reports

    importlib.reload(app_paths)
    importlib.reload(db)
    importlib.reload(ledger)
    importlib.reload(reports)
    db.configure_db_path(tmp_path / "finance.db")
    db.init_db(tmp_path / "finance.db")
    yield db, ledger, reports


def _post(ledger, summary, lines):
    ok, message, voucher_id = ledger.create_voucher(
        "2026-07-10",
        summary,
        lines,
        creator="制单员",
        source="manual",
        idempotency_key="test:" + summary,
        submit=True,
    )
    assert ok, message
    ok, message = ledger.audit_voucher(voucher_id, "审核员")
    assert ok, message
    ok, message = ledger.post_voucher(voucher_id, "记账员")
    assert ok, message
    return voucher_id


def test_books_export_all_formats(finance_modules, tmp_path):
    db, ledger, reports = finance_modules
    _post(
        ledger,
        "资本投入",
        [
            {"account_code": "1002", "debit_cents": 20_000_000, "credit_cents": 0},
            {"account_code": "3001", "debit_cents": 0, "credit_cents": 20_000_000},
        ],
    )
    _post(
        ledger,
        "销售收款",
        [
            {"account_code": "1002", "debit_cents": 5_000_000, "credit_cents": 0},
            {"account_code": "5001", "debit_cents": 0, "credit_cents": 5_000_000},
        ],
    )

    ab = reports.account_balance_table("2026-07")
    jl = reports.chronological_ledger("2026-07")
    assert jl["balanced"] is True
    cache = {
        "period": "2026-07",
        "account_balance_table": ab,
        "chronological_ledger": jl,
    }

    from gui.main import FinanceApp

    class Stub:
        pass

    Stub._resolve_cjk_font = FinanceApp._resolve_cjk_font
    stub = Stub()

    # --- CSV ---
    csv_path = tmp_path / "账簿.csv"
    FinanceApp._export_books_csv(stub, csv_path, cache)
    text = csv_path.read_text(encoding="utf-8-sig")
    assert "科目余额表" in text and "序时簿" in text
    assert "1002" in text and "银行存款" in text

    # --- Excel ---
    xlsx_path = tmp_path / "账簿.xlsx"
    FinanceApp._export_books_xlsx(stub, xlsx_path, cache)
    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path)
    assert wb.sheetnames == ["科目余额表", "序时簿"]
    assert wb["科目余额表"].max_row >= 3  # 表头 + 至少一行科目
    assert wb["序时簿"].max_row >= 2  # 表头 + 至少一条分录

    # --- PDF（含中文，验证 CJK 字体不炸） ---
    pdf_path = tmp_path / "账簿.pdf"
    FinanceApp._export_books_pdf(stub, pdf_path, cache)
    data = pdf_path.read_bytes()
    assert data[:4] == b"%PDF"
    assert len(data) > 2000
